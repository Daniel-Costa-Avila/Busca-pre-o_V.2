from __future__ import annotations

import asyncio
import hashlib
import json
import os
import re
import secrets
import shutil
import smtplib
import ssl
import subprocess
import sys
import time
import unicodedata
import requests
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from email.message import EmailMessage
from pathlib import Path
from queue import Queue
from threading import Lock, Thread
from typing import Any
from urllib.parse import urlencode

from fastapi import FastAPI, File, Form, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, FileResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook, Workbook
from pydantic import BaseModel, Field

from App.config import get_server_settings

BASE_DIR = Path(__file__).resolve().parent.parent
RUNS_DIR = BASE_DIR / "runs"
DEFAULT_INPUT = BASE_DIR / "input.xlsx"
FIXED_MODEL_TEMPLATE = Path(r"C:\Users\daniel.avila\Desktop\modelo.xlsx")
DAILY_INPUT_DIR = Path(r"C:\Users\daniel.avila\Desktop\AGENTE_DE_PRECOS\Planilha diaria")
SECONDARY_INPUT_DIR = Path(r"C:\Users\daniel.avila\Desktop\AGENTE_DE_PRECOS\Planilha diaria")
SECONDARY_ML_PATH = Path(r"C:\Users\daniel.avila\Desktop\AGENTE_DE_PRECOS\Planilha diaria\Relatorio_Financeiro.xlsx")
SECONDARY_ML_REQUIRED = True
PRIMARY_SCHEDULED_INPUT = Path(r"C:\Users\daniel.avila\Desktop\AGENTE_DE_PRECOS\input.xlsx")
DAILY_RESULT_DIR = Path(r"C:\Users\daniel.avila\Desktop\AGENTE_DE_PRECOS\Resultado Diario")
MANUAL_RESULT_DIR = Path(r"C:\Users\daniel.avila\Desktop\AGENTE_DE_PRECOS\Resultado Manual")
SERVER_SETTINGS = get_server_settings()
API_TOKEN = (os.getenv("API_TOKEN") or SERVER_SETTINGS.api_token or "").strip()
USERS_FILE = RUNS_DIR / "admin_users.json"
MARKETPLACE_FILE = RUNS_DIR / "marketplace_credentials.json"
DRIVE_SETTINGS_FILE = RUNS_DIR / "drive_settings.json"
DAILY_SCHEDULE_SETTINGS_FILE = RUNS_DIR / "daily_schedule_settings.json"
PRECOS_IMPORT_LOG_FILE = RUNS_DIR / "precos_importacoes.log"
ML_TASKS_FILE = RUNS_DIR / "ml_tasks.json"

ADMIN_SESSION_HEADER = "x-user-session"
ONLY_ONE_JOB_MESSAGE = "Ja existe uma busca em andamento. Aguarde finalizar para iniciar outra."

RESULT_SYSTEM_NAME = (os.getenv("RESULT_SYSTEM_NAME") or "Busca-Preço").strip() or "Busca-Preço"

def _env_bool(name: str, default: bool) -> bool:
    raw = os.getenv(name)
    if raw is None or not str(raw).strip():
        return default
    value = str(raw).strip().lower()
    if value in {"1", "true", "yes", "on"}:
        return True
    if value in {"0", "false", "no", "off"}:
        return False
    return default


def _env_str(name: str, default: str = "") -> str:
    raw = os.getenv(name)
    if raw is None or not str(raw).strip():
        return default.strip()
    return str(raw).strip()


def _env_int(name: str, default: int) -> int:
    raw = os.getenv(name)
    if raw is None or not str(raw).strip():
        return default
    try:
        return int(str(raw).strip())
    except ValueError:
        return default


def _env_float(name: str, default: float) -> float:
    raw = os.getenv(name)
    if raw is None or not str(raw).strip():
        return default
    try:
        return float(str(raw).strip())
    except ValueError:
        return default


def _slugify_filename(text: str) -> str:
    raw = str(text or "").strip()
    if not raw:
        return "arquivo"
    normalized = unicodedata.normalize("NFKD", raw)
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    normalized = normalized.replace(" ", "-")
    normalized = re.sub(r"[^A-Za-z0-9_-]+", "-", normalized)
    normalized = re.sub(r"[-_]{2,}", "-", normalized)
    normalized = normalized.strip("-_")
    return normalized or "arquivo"


def _daily_result_filename(dt: datetime | None = None) -> str:
    return "resultados-agentes-de-precos.xlsx"


def _daily_result_path(dt: datetime | None = None) -> Path:
    return DAILY_RESULT_DIR / _daily_result_filename(dt)


def _manual_result_filename(dt: datetime | None = None) -> str:
    stamp = (dt or datetime.now()).strftime("%Y%m%d_%H%M%S")
    return f"resultado_manual_{stamp}.xlsx"


def _manual_result_path(dt: datetime | None = None) -> Path:
    return MANUAL_RESULT_DIR / _manual_result_filename(dt)


SMTP_HOST = _env_str("SMTP_HOST", SERVER_SETTINGS.smtp_host)
SMTP_PORT = _env_int("SMTP_PORT", SERVER_SETTINGS.smtp_port)
SMTP_USER = _env_str("SMTP_USER", SERVER_SETTINGS.smtp_user)
SMTP_PASSWORD = _env_str("SMTP_PASSWORD", SERVER_SETTINGS.smtp_password)
SMTP_SENDER = _env_str("SMTP_SENDER", SERVER_SETTINGS.smtp_sender)
SMTP_USE_SSL = _env_bool("SMTP_USE_SSL", SERVER_SETTINGS.smtp_use_ssl)
SMTP_USE_TLS = _env_bool("SMTP_USE_TLS", SERVER_SETTINGS.smtp_use_tls)

GMAIL_USER = _env_str("GMAIL_USER", SERVER_SETTINGS.gmail_user or SERVER_SETTINGS.gmail_address)
GMAIL_APP_PASSWORD = _env_str("GMAIL_APP_PASSWORD", SERVER_SETTINGS.gmail_app_password)
GMAIL_SENDER = _env_str("GMAIL_SENDER", SERVER_SETTINGS.gmail_sender)

if GMAIL_USER and GMAIL_APP_PASSWORD:
    SMTP_HOST = "smtp.gmail.com"
    SMTP_PORT = 587
    SMTP_USER = GMAIL_USER
    SMTP_PASSWORD = GMAIL_APP_PASSWORD
    SMTP_SENDER = GMAIL_SENDER or GMAIL_USER
    SMTP_USE_TLS = True
    SMTP_USE_SSL = False

if not SMTP_SENDER and SMTP_USER:
    SMTP_SENDER = SMTP_USER

SMTP_RETRY_ATTEMPTS = max(1, _env_int("SMTP_RETRY_ATTEMPTS", SERVER_SETTINGS.smtp_retry_attempts))
SMTP_RETRY_MIN_SECONDS = max(0.2, _env_float("SMTP_RETRY_MIN_SECONDS", SERVER_SETTINGS.smtp_retry_min_seconds))
SMTP_RETRY_MAX_SECONDS = max(SMTP_RETRY_MIN_SECONDS, _env_float("SMTP_RETRY_MAX_SECONDS", SERVER_SETTINGS.smtp_retry_max_seconds))
EMAIL_SUBJECT_PREFIX = _env_str("EMAIL_SUBJECT_PREFIX", SERVER_SETTINGS.email_subject_prefix)
DEFAULT_EMAIL_RECIPIENTS = _env_str("EMAIL_RECIPIENTS", SERVER_SETTINGS.email_recipients)
DEFAULT_EMAIL_SEND_ON = _env_str("EMAIL_SEND_ON", SERVER_SETTINGS.email_send_on).upper()
EMAIL_DELIVERY_MODE = _env_str("EMAIL_DELIVERY_MODE", SERVER_SETTINGS.email_delivery_mode or "auto").lower()
if EMAIL_DELIVERY_MODE not in {"auto", "smtp", "outlook"}:
    EMAIL_DELIVERY_MODE = "auto"
OUTLOOK_FALLBACK_ENABLED = _env_bool("OUTLOOK_FALLBACK_ENABLED", default=False)
EMAIL_ENABLED = _env_bool("EMAIL_ENABLED", default=True)

TWILIO_ACCOUNT_SID = _env_str("TWILIO_ACCOUNT_SID", SERVER_SETTINGS.twilio_account_sid)
TWILIO_AUTH_TOKEN = _env_str("TWILIO_AUTH_TOKEN", SERVER_SETTINGS.twilio_auth_token)
TWILIO_WHATSAPP_FROM = _env_str("TWILIO_WHATSAPP_FROM", SERVER_SETTINGS.twilio_whatsapp_from)
DEFAULT_WHATSAPP_RECIPIENTS = _env_str("WHATSAPP_RECIPIENTS", SERVER_SETTINGS.whatsapp_recipients)
DEFAULT_WHATSAPP_SEND_ON = _env_str("WHATSAPP_SEND_ON", SERVER_SETTINGS.whatsapp_send_on).upper()
WHATSAPP_MESSAGE_PREFIX = _env_str("WHATSAPP_MESSAGE_PREFIX", SERVER_SETTINGS.whatsapp_message_prefix)
WHATSAPP_MEDIA_ENABLED = _env_bool("WHATSAPP_MEDIA_ENABLED", SERVER_SETTINGS.whatsapp_media_enabled)
APP_PUBLIC_BASE_URL = _env_str("APP_PUBLIC_BASE_URL", SERVER_SETTINGS.app_public_base_url)

GOOGLE_DRIVE_CREDENTIALS_FILE = _env_str("GOOGLE_DRIVE_CREDENTIALS_FILE", SERVER_SETTINGS.google_drive_credentials_file)
DEFAULT_GOOGLE_DRIVE_FOLDER_ID = _env_str("GOOGLE_DRIVE_FOLDER_ID", SERVER_SETTINGS.google_drive_folder_id)
DEFAULT_GOOGLE_DRIVE_SEND_ON = _env_str("GOOGLE_DRIVE_SEND_ON", SERVER_SETTINGS.google_drive_send_on).upper()
FIXED_DAILY_RUN_TIME = "02:00"
DEFAULT_DAILY_RUN_TIMES_RAW = _env_str("DAILY_RUN_TIMES", SERVER_SETTINGS.daily_run_times or "02:00,12:00,18:00")
DEFAULT_ADMIN_USERNAME_RAW = _env_str("ADMIN_DEFAULT_USERNAME", "admin")
DEFAULT_ADMIN_PASSWORD = _env_str("ADMIN_DEFAULT_PASSWORD", "admin123") or "admin123"
DEFAULT_ADMIN_FULL_NAME = _env_str("ADMIN_DEFAULT_FULL_NAME", "Administrador") or "Administrador"
USER_SESSION_TTL_SECONDS = max(300, _env_int("USER_SESSION_TTL_SECONDS", 43200))
ML_AUTO_ROUTE_ENABLED = False


def _normalized_default_admin_username(raw: str) -> str:
    candidate = str(raw or "").strip().lower()
    if re.fullmatch(r"[a-z0-9_.-]{3,64}", candidate):
        return candidate
    return "admin"


DEFAULT_ADMIN_USERNAME = _normalized_default_admin_username(DEFAULT_ADMIN_USERNAME_RAW)

app = FastAPI()
_cors_origins_raw = _env_str("CORS_ORIGINS", "")
_cors_origins = [origin.strip() for origin in _cors_origins_raw.split(",") if origin.strip()]
if _cors_origins:
    app.add_middleware(
        CORSMiddleware,
        allow_origins=_cors_origins,
        allow_credentials=True,
        allow_methods=["*"],
        allow_headers=["*"],
    )
templates = Jinja2Templates(directory=str(BASE_DIR / "ui" / "templates"))
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "ui" / "static")), name="static")


@dataclass
class Job:
    job_id: str
    trigger: str = "MANUAL"
    status: str = "QUEUED"
    created_at: float = field(default_factory=time.time)
    started_at: float | None = None
    finished_at: float | None = None
    input_path: Path | None = None
    output_path: Path | None = None
    error: str | None = None
    email_status: str = "PENDING"
    email_error: str | None = None
    email_recipients: list[str] = field(default_factory=list)
    auto_email_enabled: bool = False
    email_sent_at: float | None = None
    email_extra_attachment: str = ""
    whatsapp_status: str = "PENDING"
    whatsapp_error: str | None = None
    whatsapp_recipients: list[str] = field(default_factory=list)
    auto_whatsapp_enabled: bool = False
    whatsapp_sent_at: float | None = None
    drive_status: str = "PENDING"
    drive_error: str | None = None
    drive_file_id: str | None = None
    drive_file_url: str | None = None
    drive_uploaded_at: float | None = None
    drive_folder_id: str = ""
    drive_credentials_file: str = ""
    auto_drive_enabled: bool = False
    output_mode: str = "completa"


class EmailSettingsPayload(BaseModel):
    recipients: str = ""
    auto_send_manual: bool = False
    auto_send_daily: bool = False


class WhatsAppSettingsPayload(BaseModel):
    recipients: str = ""
    auto_send_manual: bool = False
    auto_send_daily: bool = False


class DriveSettingsPayload(BaseModel):
    credentials_file: str = ""
    folder_id: str = ""
    auto_send_manual: bool = False
    auto_send_daily: bool = False


class DailyScheduleSettingsPayload(BaseModel):
    run_time: str = ""
    enabled: bool | None = None
    extra_slots_enabled: bool | None = None


class AdminAuthenticatePayload(BaseModel):
    username: str
    password: str


class PrecoImportItem(BaseModel):
    sku_ref: str = ""
    titulo: str = ""
    preco: float
    moeda: str = "BRL"
    url: str = ""


class PrecosImportPayload(BaseModel):
    fonte: str = Field(default="mercado_livre", min_length=1)
    coletado_em: str = ""
    itens: list[PrecoImportItem] = Field(default_factory=list)


class MLTarefaCreatePayload(BaseModel):
    entries: list[str] = Field(default_factory=list)
    limite_por_pesquisa: int = 10
    delay_ms: int = 1200


class MLTarefaResultadoPayload(BaseModel):
    status: str = Field(default="done", min_length=1)
    itens: list[dict[str, Any]] = Field(default_factory=list)
    erro: str = ""


class AdminCreateUserPayload(BaseModel):
    username: str
    password: str
    full_name: str | None = None
    role: str = "operator"
    is_active: bool = True


class AdminUpdateUserPayload(BaseModel):
    full_name: str | None = None
    role: str | None = None
    is_active: bool | None = None
    new_password: str | None = None


class MagaluCredentialsPayload(BaseModel):
    api_key: str | None = None
    api_key_id: str | None = None
    api_key_secret: str | None = None
    access_token: str | None = None
    refresh_token: str | None = None
    api_base: str | None = None
    token_url: str | None = None
    prices_path: str | None = None
    seller_id: str | None = None
    redirect_uri: str | None = None
    auth_scope: str | None = None


jobs: dict[str, Job] = {}
job_queue: Queue[str] = Queue()
latest_manual_job_id: str | None = None
latest_daily_job_id: str | None = None

_runner_lock = Lock()
_current_process: subprocess.Popen[str] | None = None
_current_running_job_id: str | None = None
_stop_requested_jobs: set[str] = set()
_daily_scheduler_lock = Lock()
_daily_triggered_slots: set[str] = set()

_email_settings_lock = Lock()
_runtime_email_recipients: list[str] = []
_runtime_auto_email_manual = False
_runtime_auto_email_daily = False

ALWAYS_EMAIL_RECIPIENTS = [
    "daniel.avila@colchoes.ind.br",
    "tdefrete@gmail.com",
]
FEATURE_AVISTA_ENABLED = False

_whatsapp_settings_lock = Lock()
_runtime_whatsapp_recipients: list[str] = []
_runtime_auto_whatsapp_manual = False
_runtime_auto_whatsapp_daily = False

_drive_settings_lock = Lock()
_runtime_drive_credentials_file = ""
_runtime_drive_folder_id = ""
_runtime_auto_drive_manual = False
_runtime_auto_drive_daily = False

_daily_settings_lock = Lock()
_runtime_daily_slot_states: dict[str, bool] = {}
_runtime_daily_custom_times: list[str] = []

_admin_users_lock = Lock()
_admin_users: dict[str, dict[str, Any]] = {}
_admin_sessions_lock = Lock()
_admin_sessions: dict[str, dict[str, Any]] = {}

_marketplace_lock = Lock()
_precos_import_log_lock = Lock()
_ml_tasks_lock = Lock()
_marketplace_credentials: dict[str, dict[str, str]] = {}
_ml_tasks: dict[str, dict[str, Any]] = {}


def _now_ts() -> float:
    return time.time()


def _parse_bool(raw: str | None, default: bool | None = None) -> bool | None:
    if raw is None:
        return default
    v = str(raw).strip().lower()
    if v in {"1", "true", "yes", "on"}:
        return True
    if v in {"0", "false", "no", "off"}:
        return False
    return default


def _parse_recipients(raw: str | None) -> list[str]:
    if not raw:
        return []
    out: list[str] = []
    seen: set[str] = set()
    for part in re.split(r"[,;\n]+", raw):
        email = part.strip()
        if not email or "@" not in email or "." not in email:
            continue
        key = email.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(email)
    return out


def _merge_recipients(primary: list[str], extra: list[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for item in primary + extra:
        value = str(item or "").strip()
        if not value:
            continue
        key = value.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(value)
    return out


def _normalize_whatsapp_recipient(raw: str | None) -> str | None:
    value = str(raw or "").strip()
    if not value:
        return None

    if value.lower().startswith("whatsapp:"):
        value = value[len("whatsapp:") :].strip()

    cleaned = re.sub(r"[^\d+]", "", value)
    if cleaned.startswith("00"):
        cleaned = f"+{cleaned[2:]}"
    if cleaned and not cleaned.startswith("+"):
        cleaned = f"+{cleaned}"

    digits = re.sub(r"\D", "", cleaned)
    if len(digits) < 10:
        return None
    return f"whatsapp:{cleaned}"


def _parse_whatsapp_recipients(raw: str | None) -> list[str]:
    if not raw:
        return []
    out: list[str] = []
    seen: set[str] = set()
    for part in re.split(r"[,;\n]+", raw):
        recipient = _normalize_whatsapp_recipient(part)
        if not recipient:
            continue
        key = recipient.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(recipient)
    return out


def _recipients_to_csv(values: list[str]) -> str:
    return ", ".join(values)


def _normalize_drive_folder_id(raw: str | None) -> str:
    value = str(raw or "").strip()
    if not value:
        return ""
    match = re.search(r"/folders/([A-Za-z0-9_-]{10,})", value)
    if match:
        return match.group(1)
    if re.fullmatch(r"[A-Za-z0-9_-]{10,}", value):
        return value
    return ""


def _normalize_drive_credentials_file(raw: str | None) -> str:
    value = str(raw or "").strip().strip('"').strip("'")
    if not value:
        return ""
    return str(Path(value).expanduser())


def _normalize_email_attachment_path(raw: str | None) -> str:
    value = str(raw or "").strip().strip('"').strip("'")
    if not value:
        return ""
    candidate = Path(value).expanduser()
    if not candidate.is_absolute():
        candidate = BASE_DIR / candidate
    try:
        resolved = candidate.resolve()
    except Exception:
        resolved = candidate
    if resolved.exists() and resolved.is_file():
        return str(resolved)
    return ""


def _latest_xlsx_in_dir(dir_path: Path, exclude_names: set[str] | None = None) -> Path | None:
    if not dir_path.exists() or not dir_path.is_dir():
        return None
    exclude = {name.lower() for name in (exclude_names or set())}
    files = [
        p
        for p in dir_path.iterdir()
        if p.is_file()
        and p.suffix.lower() == ".xlsx"
        and p.name.lower() not in exclude
    ]
    if not files:
        return None
    files.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return files[0]


def _resolve_secondary_attachment() -> Path | None:
    if SECONDARY_ML_PATH.exists() and SECONDARY_ML_PATH.is_file():
        return SECONDARY_ML_PATH
    return _latest_xlsx_in_dir(
        SECONDARY_INPUT_DIR,
        exclude_names={"resultados-agentes-de-precos.xlsx"},
    )


def _archive_result_by_origin(job: Job) -> None:
    if not job.output_path or not job.output_path.exists():
        return
    try:
        now = datetime.now()
        if str(job.trigger or "").upper() == "MANUAL":
            MANUAL_RESULT_DIR.mkdir(parents=True, exist_ok=True)
            target = _manual_result_path(now)
        else:
            DAILY_RESULT_DIR.mkdir(parents=True, exist_ok=True)
            target = _daily_result_path(now)
        shutil.copy2(job.output_path, target)
    except Exception:
        # nao interrompe o job se falhar
        return


def _normalize_text(value: object) -> str:
    text = str(value or "").strip().lower()
    if not text:
        return ""
    normalized = unicodedata.normalize("NFKD", text)
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", normalized)


def _to_float(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    if not text:
        return None

    match_brl = re.search(r"R\$\s*([\d\.]+,\d{2})", text, re.IGNORECASE)
    if match_brl:
        raw = match_brl.group(1).replace(".", "").replace(",", ".")
        try:
            return float(raw)
        except ValueError:
            pass

    cleaned = re.sub(r"[^\d,.\-]", "", text)
    if not cleaned:
        return None

    if "," in cleaned and "." in cleaned:
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")
    elif "," in cleaned:
        cleaned = cleaned.replace(",", ".")

    try:
        return float(cleaned)
    except ValueError:
        return None


def _to_number(value) -> float | int | None:
    num = _to_float(value)
    if num is None:
        return None
    if float(num).is_integer():
        return int(num)
    return num


def _extract_url_from_cell(cell) -> str:
    def _extract_url_from_text(raw: object) -> str:
        text = str(raw or "").strip()
        if not text:
            return ""

        formula_match = re.search(r'HYPERLINK\(\s*"([^"]+)"', text, re.IGNORECASE)
        if formula_match:
            return formula_match.group(1).strip()

        inline_match = re.search(r"https?://[^\s\"'<>]+", text, re.IGNORECASE)
        if inline_match:
            return inline_match.group(0).rstrip(".,;)")

        return ""

    visible_url = _extract_url_from_text(cell.value)
    if visible_url:
        return visible_url

    if getattr(cell, "hyperlink", None) and cell.hyperlink.target:
        return str(cell.hyperlink.target).strip()

    return ""


def _find_output_col(headers: list[str], names: set[str]) -> int | None:
    normalized_names = {_normalize_text(name) for name in names if _normalize_text(name)}
    if not normalized_names:
        return None
    for idx, header in enumerate(headers):
        if _normalize_text(header) in normalized_names:
            return idx
    return None


def _build_simple_output(base_path: Path, mode: str) -> tuple[Path | None, str | None]:
    if not base_path.exists():
        return None, "Arquivo base nao encontrado."
    normalized_mode = str(mode or "").strip().lower()
    if normalized_mode not in {"a_prazo", "a_vista"}:
        normalized_mode = "a_prazo"
    if normalized_mode == "a_vista" and not FEATURE_AVISTA_ENABLED:
        normalized_mode = "a_prazo"
    try:
        wb = load_workbook(base_path)
    except Exception as exc:
        return None, f"Falha ao abrir planilha base: {type(exc).__name__}: {exc}"

    sheet_name = "Output"
    headers = ["id no Canal", "CODIGO INTERNO", "Canal", "Titulo", "Preco"] if normalized_mode == "a_prazo" else [
        "id no Canal",
        "CODIGO INTERNO",
        "Canal",
        "Titulo",
        "A vista",
    ]

    ws_out = wb["Output"] if "Output" in wb.sheetnames else wb.active
    output_headers = [
        str(ws_out.cell(row=1, column=col).value or "").strip()
        for col in range(1, ws_out.max_column + 1)
    ]
    out_id = _find_output_col(output_headers, {
        "id no canal",
        "id_no_canal",
        "id canal",
        "id_canal",
        "codigo lojista",
        "codigo_lojista",
        "sku canal",
        "sku",
    })
    out_codigo = _find_output_col(output_headers, {
        "codigo interno",
        "codigo_interno",
        "id_produto",
        "produto",
        "id",
        "codigo",
        "codigo_produto",
        "product_id",
    })
    out_canal = _find_output_col(output_headers, {"canal", "loja", "marketplace", "seller"})
    out_titulo = _find_output_col(output_headers, {"titulo", "tÃ­tulo", "nome", "descricao", "descriÃ§Ã£o"})
    out_prazo = _find_output_col(output_headers, {"a prazo", "a_prazo", "prazo", "preco"})
    out_avista = _find_output_col(output_headers, {"a vista", "a_vista", "avista"})

    suffix = "aprazo" if normalized_mode == "a_prazo" else "avista"
    simple_path = base_path.parent / f"{base_path.stem}_{suffix}.xlsx"
    try:
        simple_only = Workbook()
        target_ws = simple_only.active
        target_ws.title = sheet_name
        target_ws.append(headers)
        for row in range(2, ws_out.max_row + 1):
            target_ws.append([
                _to_number(ws_out.cell(row=row, column=out_id + 1).value) if out_id is not None else None,
                _to_number(ws_out.cell(row=row, column=out_codigo + 1).value) if out_codigo is not None else None,
                ws_out.cell(row=row, column=out_canal + 1).value if out_canal is not None else None,
                ws_out.cell(row=row, column=out_titulo + 1).value if out_titulo is not None else None,
                _to_number(ws_out.cell(row=row, column=(out_prazo + 1) if normalized_mode == "a_prazo" else (out_avista + 1)).value)
                if (out_prazo is not None and normalized_mode == "a_prazo") or (out_avista is not None and normalized_mode == "a_vista")
                else None,
            ])
        simple_only.save(simple_path)
    except Exception as exc:
        label = "a prazo" if normalized_mode == "a_prazo" else "a vista"
        return None, f"Falha ao salvar planilha {label}: {type(exc).__name__}: {exc}"

    return simple_path, None


def _build_merged_output(job: Job) -> tuple[Path | None, str | None]:
    if not job.output_path or not job.output_path.exists():
        return None, "Arquivo de output nao encontrado."
    extra_attachment = _normalize_email_attachment_path(job.email_extra_attachment)
    if not extra_attachment:
        return None, "Anexo extra nao encontrado."

    try:
        wb_main = load_workbook(job.output_path)
        ws_main = wb_main["Output"] if "Output" in wb_main.sheetnames else wb_main.active
    except Exception as exc:
        return None, f"Falha ao abrir output: {type(exc).__name__}: {exc}"

    try:
        wb_ml = load_workbook(extra_attachment)
        ws_ml = wb_ml.active
    except Exception as exc:
        return None, f"Falha ao abrir planilha Mercado Livre: {type(exc).__name__}: {exc}"

    headers = [
        str(ws_main.cell(row=1, column=col).value or "").strip()
        for col in range(1, ws_main.max_column + 1)
    ]
    total_cols = max(1, len(headers))

    out_id = _find_output_col(headers, {
        "id no canal",
        "id_no_canal",
        "id canal",
        "id_canal",
        "codigo lojista",
        "codigo_lojista",
        "sku canal",
        "sku",
    })

    seen_ids: set[str] = set()
    merged_wb = Workbook()
    merged_ws = merged_wb.active
    merged_ws.title = "Output"
    merged_ws.append(headers)
    for row in ws_main.iter_rows(min_row=2, values_only=True):
        if out_id is not None and out_id < len(row):
            raw_id = row[out_id]
            key = str(raw_id or "").strip()
            if key:
                if key in seen_ids:
                    continue
                seen_ids.add(key)
        merged_ws.append(list(row))

    out_codigo = _find_output_col(headers, {
        "codigo interno",
        "codigo_interno",
        "id_produto",
        "produto",
        "id",
        "codigo",
        "codigo_produto",
        "product_id",
    })
    out_canal = _find_output_col(headers, {"canal", "loja", "marketplace", "seller"})
    out_titulo = _find_output_col(headers, {"titulo", "tÃ­tulo", "nome", "descricao", "descriÃ§Ã£o"})
    out_prazo = _find_output_col(headers, {"a prazo", "a_prazo", "prazo", "preco"})
    out_avista = _find_output_col(headers, {"a vista", "a_vista", "avista"})
    out_parcelamento = _find_output_col(headers, {"parcelamento"})
    out_link = _find_output_col(headers, {"link", "url", "href"})
    out_status = _find_output_col(headers, {"status"})
    out_data = _find_output_col(headers, {"data_pesquisa", "data pesquisa", "data"})
    out_hora = _find_output_col(headers, {"hora_pesquisa", "hora pesquisa", "hora"})
    out_origem_execucao = _find_output_col(headers, {"origem_execucao", "origem execucao"})
    out_origem_dados = _find_output_col(headers, {"origem_dados", "origem dados", "source"})

    expected_headers = {"id no canal", "codigo interno", "canal", "titulo", "prazo"}
    first_row_values = [
        _normalize_text(ws_ml.cell(row=1, column=c).value)
        for c in range(1, 6)
    ]
    has_header = any(value in expected_headers for value in first_row_values)
    start_row = 2 if has_header else 1

    for row in range(start_row, ws_ml.max_row + 1):
        id_no_canal = ws_ml.cell(row=row, column=1).value
        codigo_interno = ws_ml.cell(row=row, column=2).value
        canal = ws_ml.cell(row=row, column=3).value
        titulo = ws_ml.cell(row=row, column=4).value
        a_prazo = ws_ml.cell(row=row, column=5).value
        link = _extract_url_from_cell(ws_ml.cell(row=row, column=6))

        if not any([id_no_canal, codigo_interno, canal, titulo, a_prazo, link]):
            continue
        if out_id is not None:
            key = str(id_no_canal or "").strip()
            if key:
                if key in seen_ids:
                    continue
                seen_ids.add(key)

        out_row = [None] * total_cols
        if out_id is not None:
            out_row[out_id] = id_no_canal
        if out_codigo is not None:
            out_row[out_codigo] = codigo_interno
        if out_canal is not None:
            out_row[out_canal] = canal
        if out_titulo is not None:
            out_row[out_titulo] = titulo
        if out_prazo is not None:
            out_row[out_prazo] = a_prazo
        if out_avista is not None:
            out_row[out_avista] = None
        if out_parcelamento is not None:
            out_row[out_parcelamento] = None
        if out_link is not None:
            out_row[out_link] = link
        if out_status is not None:
            out_row[out_status] = "IMPORTADO"
        if out_origem_execucao is not None:
            out_row[out_origem_execucao] = "IMPORTADO"
        if out_origem_dados is not None:
            out_row[out_origem_dados] = "mercado_livre"
        if out_data is not None:
            out_row[out_data] = None
        if out_hora is not None:
            out_row[out_hora] = None

        merged_ws.append(out_row)

    merged_path = job.output_path.parent / f"{job.output_path.stem}_merged.xlsx"
    try:
        merged_wb.save(merged_path)
    except Exception as exc:
        return None, f"Falha ao salvar planilha unificada: {type(exc).__name__}: {exc}"

    return merged_path, None


def _drive_credentials_exists(path_value: str | None) -> bool:
    path = _normalize_drive_credentials_file(path_value)
    if not path:
        return False
    try:
        return Path(path).is_file()
    except Exception:
        return False


def _extract_drive_service_account_email(path_value: str | None) -> str | None:
    path = _normalize_drive_credentials_file(path_value)
    if not path:
        return None
    payload = _read_json(Path(path), {})
    if not isinstance(payload, dict):
        return None
    email = str(payload.get("client_email") or "").strip()
    return email or None


def _read_json(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _counts() -> dict[str, int]:
    out = {"QUEUED": 0, "RUNNING": 0, "DONE": 0, "FAILED": 0, "STOPPED": 0}
    for job in jobs.values():
        if job.status in out:
            out[job.status] += 1
    return out


def _new_job_id(trigger: str = "MANUAL") -> str:
    """
    Gera identificador baseado em data/hora da pesquisa.
    Formato base: YYYYMMDD_HHMMSS.
    Se houver colisao no mesmo segundo, adiciona sufixo incremental.
    """
    base = datetime.now().strftime("%Y%m%d_%H%M%S")
    candidate = base
    suffix = 1
    while candidate in jobs or (RUNS_DIR / candidate).exists():
        candidate = f"{base}_{suffix:02d}"
        suffix += 1
    return candidate


def _output_filename(job_id: str) -> str:
    return f"pesquisa_{job_id}.xlsx"


def _job_payload(job: Job | None) -> dict | None:
    if job is None:
        return None
    output_available = bool(job.output_path and job.output_path.exists())
    return {
        "job_id": job.job_id,
        "trigger": job.trigger,
        "status": job.status,
        "created_at": job.created_at,
        "started_at": job.started_at,
        "finished_at": job.finished_at,
        "error": job.error,
        "email_status": job.email_status,
        "email_error": job.email_error,
        "email_sent_at": job.email_sent_at,
        "whatsapp_status": job.whatsapp_status,
        "whatsapp_error": job.whatsapp_error,
        "whatsapp_sent_at": job.whatsapp_sent_at,
        "drive_status": job.drive_status,
        "drive_error": job.drive_error,
        "drive_uploaded_at": job.drive_uploaded_at,
        "drive_file_id": job.drive_file_id,
        "drive_file_url": job.drive_file_url,
        "output_available": output_available,
        "auto_email_enabled": job.auto_email_enabled,
        "auto_whatsapp_enabled": job.auto_whatsapp_enabled,
        "auto_drive_enabled": job.auto_drive_enabled,
        "output_mode": job.output_mode,
        "email_recipients_csv": _recipients_to_csv(job.email_recipients),
        "email_extra_attachment": job.email_extra_attachment,
        "whatsapp_recipients_csv": _recipients_to_csv(job.whatsapp_recipients),
        "drive_folder_id": job.drive_folder_id,
    }


def _latest_manual_job() -> Job | None:
    global latest_manual_job_id
    if latest_manual_job_id and latest_manual_job_id in jobs:
        return jobs[latest_manual_job_id]
    manual = [j for j in jobs.values() if j.trigger == "MANUAL"]
    if not manual:
        return None
    manual.sort(key=lambda item: item.created_at, reverse=True)
    latest_manual_job_id = manual[0].job_id
    return manual[0]


def _latest_daily_job() -> Job | None:
    global latest_daily_job_id
    if latest_daily_job_id and latest_daily_job_id in jobs:
        return jobs[latest_daily_job_id]
    daily = [j for j in jobs.values() if j.trigger != "MANUAL"]
    if not daily:
        return None
    daily.sort(key=lambda item: item.created_at, reverse=True)
    latest_daily_job_id = daily[0].job_id
    return daily[0]


def _latest_output_job() -> Job | None:
    candidates = [j for j in jobs.values() if j.output_path and j.output_path.exists()]
    if not candidates:
        return None
    candidates.sort(key=lambda item: (item.finished_at or 0, item.created_at), reverse=True)
    return candidates[0]


def _latest_done_job() -> Job | None:
    candidates = [j for j in jobs.values() if j.status == "DONE" and j.output_path and j.output_path.exists()]
    if not candidates:
        return None
    candidates.sort(key=lambda item: (item.finished_at or 0, item.created_at), reverse=True)
    return candidates[0]


def _default_magalu_credentials() -> dict[str, str]:
    return {
        "api_key": "",
        "api_key_id": "",
        "api_key_secret": "",
        "access_token": "",
        "refresh_token": "",
        "api_base": "https://api.magalu.com",
        "token_url": "https://id.magalu.com/oauth/token",
        "prices_path": "/seller/v1/portfolios/prices/{sku}",
        "seller_id": "",
        "redirect_uri": "",
        "auth_scope": "openid profile email",
        "oauth_state": "",
    }


def _load_marketplace_credentials() -> None:
    global _marketplace_credentials
    raw = _read_json(MARKETPLACE_FILE, {})
    magalu = _default_magalu_credentials()
    if isinstance(raw, dict) and isinstance(raw.get("magalu"), dict):
        for key in magalu:
            value = raw["magalu"].get(key)
            if value is not None:
                magalu[key] = str(value).strip()
    _marketplace_credentials = {"magalu": magalu}


def _save_marketplace_credentials() -> None:
    _write_json(MARKETPLACE_FILE, _marketplace_credentials)


def _hash_password(password: str, salt_hex: str) -> str:
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), bytes.fromhex(salt_hex), 180_000)
    return digest.hex()


def _new_password(password: str) -> tuple[str, str]:
    salt_hex = secrets.token_hex(16)
    return _hash_password(password, salt_hex), salt_hex


def _public_user(record: dict[str, Any]) -> dict[str, Any]:
    return {
        "username": record.get("username"),
        "full_name": record.get("full_name"),
        "role": record.get("role"),
        "is_active": bool(record.get("is_active", True)),
        "created_at": datetime.fromtimestamp(float(record.get("created_at") or _now_ts())).isoformat(timespec="seconds"),
        "updated_at": datetime.fromtimestamp(float(record.get("updated_at") or _now_ts())).isoformat(timespec="seconds"),
    }


def _load_users() -> None:
    global _admin_users
    users = _read_json(USERS_FILE, {})
    if not isinstance(users, dict):
        users = {}
    normalized: dict[str, dict[str, Any]] = {}
    for username, payload in users.items():
        if not isinstance(payload, dict):
            continue
        key = str(username).strip().lower()
        if not key:
            continue
        normalized[key] = payload

    has_active_admin = any(
        str(payload.get("role") or "").strip().lower() == "admin" and bool(payload.get("is_active", True))
        for payload in normalized.values()
        if isinstance(payload, dict)
    )

    if not has_active_admin:
        hash_value, salt_hex = _new_password(DEFAULT_ADMIN_PASSWORD)
        now = _now_ts()
        normalized[DEFAULT_ADMIN_USERNAME] = {
            "username": DEFAULT_ADMIN_USERNAME,
            "full_name": DEFAULT_ADMIN_FULL_NAME,
            "role": "admin",
            "is_active": True,
            "password_hash": hash_value,
            "salt": salt_hex,
            "created_at": now,
            "updated_at": now,
        }
    _admin_users = normalized
    _write_json(USERS_FILE, _admin_users)


def _admin_actor(request: Request, require_admin: bool = True) -> tuple[dict[str, Any] | None, JSONResponse | None]:
    token = str(request.headers.get(ADMIN_SESSION_HEADER) or "").strip()
    if not token:
        return None, JSONResponse({"error": "Sessao admin ausente."}, status_code=401)
    with _admin_sessions_lock:
        session = _admin_sessions.get(token)
    if not session:
        return None, JSONResponse({"error": "Sessao admin invalida."}, status_code=401)
    if float(session.get("expires_at") or 0) <= _now_ts():
        with _admin_sessions_lock:
            _admin_sessions.pop(token, None)
        return None, JSONResponse({"error": "Sessao admin expirada."}, status_code=401)

    username = str(session.get("username") or "").strip().lower()
    with _admin_users_lock:
        actor = _admin_users.get(username)
    if not actor or not actor.get("is_active"):
        return None, JSONResponse({"error": "Usuario inativo."}, status_code=403)
    if require_admin and str(actor.get("role") or "").lower() != "admin":
        return None, JSONResponse({"error": "Permissao negada."}, status_code=403)
    return actor, None


def _load_email_settings() -> None:
    global _runtime_email_recipients, _runtime_auto_email_manual, _runtime_auto_email_daily
    with _email_settings_lock:
        recipients_raw = os.getenv("EMAIL_RECIPIENTS", "") or os.getenv("EMAIL_TO", "") or DEFAULT_EMAIL_RECIPIENTS
        _runtime_email_recipients = _parse_recipients(recipients_raw)
        send_on = str(os.getenv("EMAIL_SEND_ON", "") or DEFAULT_EMAIL_SEND_ON).upper()
        _runtime_auto_email_manual = "MANUAL" in send_on
        _runtime_auto_email_daily = "AUTO_DIARIO" in send_on


def _load_whatsapp_settings() -> None:
    global _runtime_whatsapp_recipients, _runtime_auto_whatsapp_manual, _runtime_auto_whatsapp_daily
    with _whatsapp_settings_lock:
        _runtime_whatsapp_recipients = _parse_whatsapp_recipients(DEFAULT_WHATSAPP_RECIPIENTS)
        send_on = DEFAULT_WHATSAPP_SEND_ON
        _runtime_auto_whatsapp_manual = "MANUAL" in send_on
        _runtime_auto_whatsapp_daily = "AUTO_DIARIO" in send_on


def _save_drive_settings() -> None:
    with _drive_settings_lock:
        payload = {
            "credentials_file": _runtime_drive_credentials_file,
            "folder_id": _runtime_drive_folder_id,
            "auto_send_manual": bool(_runtime_auto_drive_manual),
            "auto_send_daily": bool(_runtime_auto_drive_daily),
        }
    _write_json(DRIVE_SETTINGS_FILE, payload)


def _load_drive_settings() -> None:
    global _runtime_drive_credentials_file, _runtime_drive_folder_id, _runtime_auto_drive_manual, _runtime_auto_drive_daily
    saved = _read_json(DRIVE_SETTINGS_FILE, {})
    saved_credentials = ""
    saved_folder_id = ""
    saved_auto_manual = False
    saved_auto_daily = False
    if isinstance(saved, dict):
        saved_credentials = _normalize_drive_credentials_file(saved.get("credentials_file"))
        saved_folder_id = _normalize_drive_folder_id(saved.get("folder_id"))
        saved_auto_manual = bool(saved.get("auto_send_manual"))
        saved_auto_daily = bool(saved.get("auto_send_daily"))

    with _drive_settings_lock:
        _runtime_drive_credentials_file = _normalize_drive_credentials_file(GOOGLE_DRIVE_CREDENTIALS_FILE or saved_credentials)
        _runtime_drive_folder_id = _normalize_drive_folder_id(DEFAULT_GOOGLE_DRIVE_FOLDER_ID or saved_folder_id)
        if DEFAULT_GOOGLE_DRIVE_SEND_ON:
            _runtime_auto_drive_manual = "MANUAL" in DEFAULT_GOOGLE_DRIVE_SEND_ON
            _runtime_auto_drive_daily = "AUTO_DIARIO" in DEFAULT_GOOGLE_DRIVE_SEND_ON
        else:
            _runtime_auto_drive_manual = saved_auto_manual
            _runtime_auto_drive_daily = saved_auto_daily


def _save_daily_schedule_settings() -> None:
    with _daily_settings_lock:
        payload = {
            "slot_states": dict(_runtime_daily_slot_states),
            "custom_times": list(_runtime_daily_custom_times),
        }
    _write_json(DAILY_SCHEDULE_SETTINGS_FILE, payload)


def _load_daily_schedule_settings() -> None:
    global _runtime_daily_slot_states, _runtime_daily_custom_times
    saved = _read_json(DAILY_SCHEDULE_SETTINGS_FILE, {})
    configured_slots = _extra_daily_run_times_text()
    slot_states = {run_time: True for run_time in configured_slots}
    saved_enabled = None
    saved_states = {}
    if isinstance(saved, dict):
        if "extra_slots_enabled" in saved:
            saved_enabled = bool(saved.get("extra_slots_enabled"))
        raw_states = saved.get("slot_states")
        if isinstance(raw_states, dict):
            saved_states = raw_states
        raw_custom = saved.get("custom_times")
        if isinstance(raw_custom, list):
            _runtime_daily_custom_times = [str(v).strip() for v in raw_custom if str(v).strip()]
    if saved_enabled is not None:
        for run_time in configured_slots:
            slot_states[run_time] = bool(saved_enabled)
    for run_time, enabled in saved_states.items():
        key = str(run_time).strip()
        if key in slot_states:
            slot_states[key] = bool(enabled)
    env_enabled = _parse_bool(os.getenv("DAILY_EXTRA_SLOTS_ENABLED"), default=None)
    with _daily_settings_lock:
        if env_enabled is not None:
            for run_time in configured_slots:
                slot_states[run_time] = bool(env_enabled)
        _runtime_daily_slot_states = slot_states


def _public_email_settings() -> dict[str, Any]:
    with _email_settings_lock:
        recipients = list(_runtime_email_recipients)
        auto_manual = bool(_runtime_auto_email_manual)
        auto_daily = bool(_runtime_auto_email_daily)
    return {
        "enabled": bool(EMAIL_ENABLED),
        "smtp_ready": bool(SMTP_HOST and SMTP_SENDER),
        "email_delivery_mode": EMAIL_DELIVERY_MODE,
        "outlook_fallback_enabled": bool(EMAIL_DELIVERY_MODE == "outlook" or (EMAIL_DELIVERY_MODE == "auto" and OUTLOOK_FALLBACK_ENABLED)),
        "recipients_csv": _recipients_to_csv(recipients),
        "auto_send_manual": auto_manual,
        "auto_send_daily": auto_daily,
    }


def _public_whatsapp_settings() -> dict[str, Any]:
    with _whatsapp_settings_lock:
        recipients = list(_runtime_whatsapp_recipients)
        auto_manual = bool(_runtime_auto_whatsapp_manual)
        auto_daily = bool(_runtime_auto_whatsapp_daily)
    return {
        "twilio_ready": bool(TWILIO_ACCOUNT_SID and TWILIO_AUTH_TOKEN and _normalize_whatsapp_recipient(TWILIO_WHATSAPP_FROM)),
        "from_number": _normalize_whatsapp_recipient(TWILIO_WHATSAPP_FROM),
        "recipients_csv": _recipients_to_csv(recipients),
        "auto_send_manual": auto_manual,
        "auto_send_daily": auto_daily,
        "media_enabled": bool(WHATSAPP_MEDIA_ENABLED),
        "public_base_url": APP_PUBLIC_BASE_URL,
    }


def _public_drive_settings() -> dict[str, Any]:
    with _drive_settings_lock:
        credentials_file = _runtime_drive_credentials_file
        folder_id = _runtime_drive_folder_id
        auto_manual = bool(_runtime_auto_drive_manual)
        auto_daily = bool(_runtime_auto_drive_daily)
    credentials_exists = _drive_credentials_exists(credentials_file)
    service_account_email = _extract_drive_service_account_email(credentials_file)
    return {
        "drive_ready": bool(credentials_exists and folder_id),
        "credentials_file": credentials_file,
        "credentials_exists": credentials_exists,
        "folder_id": folder_id,
        "auto_send_manual": auto_manual,
        "auto_send_daily": auto_daily,
        "service_account_email": service_account_email,
    }


def _parse_daily_run_times(raw: str | None) -> list[tuple[int, int]]:
    out: list[tuple[int, int]] = []
    for token in str(raw or "").split(","):
        value = token.strip()
        if not value:
            continue
        try:
            hour_raw, minute_raw = value.split(":", 1)
            hour = int(hour_raw)
            minute = int(minute_raw)
        except Exception:
            continue
        if not (0 <= hour <= 23 and 0 <= minute <= 59):
            continue
        out.append((hour, minute))
    return sorted(set(out))


def _format_daily_run_time(value: tuple[int, int]) -> str:
    return f"{value[0]:02d}:{value[1]:02d}"


def _fixed_daily_run_slot() -> tuple[int, int]:
    parsed = _parse_daily_run_times(FIXED_DAILY_RUN_TIME)
    return parsed[0] if parsed else (6, 0)


def _fixed_daily_run_time_text() -> str:
    return _format_daily_run_time(_fixed_daily_run_slot())


def _extra_daily_run_times() -> list[tuple[int, int]]:
    fixed = _fixed_daily_run_slot()
    all_times = _parse_daily_run_times(DEFAULT_DAILY_RUN_TIMES_RAW)
    if not all_times:
        all_times = [fixed]
    custom_times = _parse_daily_run_times(",".join(_runtime_daily_custom_times))
    merged = list({value for value in (all_times + custom_times) if value != fixed})
    return sorted(merged)


def _extra_daily_run_times_text() -> list[str]:
    return [_format_daily_run_time(value) for value in _extra_daily_run_times()]


def _daily_slot_enabled(run_time: str) -> bool:
    key = str(run_time).strip()
    if not key:
        return False
    if key == _fixed_daily_run_time_text():
        return True
    with _daily_settings_lock:
        return bool(_runtime_daily_slot_states.get(key, True))


def _set_daily_slot_enabled(run_time: str, enabled: bool) -> bool:
    key = str(run_time).strip()
    if key not in _extra_daily_run_times_text():
        return False
    with _daily_settings_lock:
        _runtime_daily_slot_states[key] = bool(enabled)
    return True


def _add_daily_custom_time(run_time: str) -> bool:
    parsed = _parse_daily_run_times(run_time)
    if not parsed:
        return False
    normalized = _format_daily_run_time(parsed[0])
    with _daily_settings_lock:
        if normalized not in _runtime_daily_custom_times:
            _runtime_daily_custom_times.append(normalized)
    return True


def _set_all_extra_daily_slots(enabled: bool) -> None:
    with _daily_settings_lock:
        for run_time in _extra_daily_run_times_text():
            _runtime_daily_slot_states[run_time] = bool(enabled)


def _daily_slots_payload(next_run: datetime | None = None) -> list[dict[str, Any]]:
    next_time = next_run.strftime("%H:%M") if next_run else ""
    fixed_time = _fixed_daily_run_time_text()
    slots: list[dict[str, Any]] = [
        {
            "time": fixed_time,
            "enabled": True,
            "fixed": True,
            "locked": True,
            "is_next": next_time == fixed_time,
            "title": "Rotina fixa da manha",
            "description": "Horario principal, sempre ativo.",
        }
    ]
    for run_time in _extra_daily_run_times_text():
        enabled = _daily_slot_enabled(run_time)
        slots.append(
            {
                "time": run_time,
                "enabled": enabled,
                "fixed": False,
                "locked": False,
                "is_next": next_time == run_time,
                "title": "Horario adicional",
                "description": "Ative ou desative conforme a rotina do dia.",
            }
        )
    return slots


def _public_daily_schedule() -> dict[str, Any]:
    fixed_run_time = _fixed_daily_run_time_text()
    effective_run_times = _daily_run_times_text()
    extra_run_times = _extra_daily_run_times_text()
    next_run = _next_daily_run_dt() if _daily_auto_enabled() else None
    slots = _daily_slots_payload(next_run)
    extra_enabled_count = sum(1 for slot in slots if not slot["fixed"] and slot["enabled"])
    return {
        "enabled": _daily_auto_enabled(),
        "fixed_run_time": fixed_run_time,
        "extra_run_times": extra_run_times,
        "extra_slots_enabled": bool(extra_enabled_count == len(extra_run_times)) if extra_run_times else False,
        "has_extra_slots": bool(extra_run_times),
        "run_times": effective_run_times,
        "configured_run_times": [slot["time"] for slot in slots],
        "slots": slots,
        "active_slot_count": len(effective_run_times),
        "extra_enabled_count": extra_enabled_count,
        "extra_total_count": len(extra_run_times),
        "next_run_time": next_run.strftime("%H:%M") if next_run else "",
        "run_hour": int(fixed_run_time.split(":")[0]),
        "run_minute": int(fixed_run_time.split(":")[1]),
        "next_run_at": next_run.timestamp() if next_run else None,
        "next_run_at_iso": next_run.isoformat(timespec="seconds") if next_run else "",
    }


def _is_public_path(path: str) -> bool:
    return path == "/api/health" or path.startswith("/static") or path == "/favicon.ico"


def _extract_token(request: Request) -> str:
    auth = request.headers.get("authorization", "")
    if auth.lower().startswith("bearer "):
        return auth[7:].strip()
    query_token = request.query_params.get("api_token")
    return (query_token or "").strip()


def _assets_version() -> str:
    """Usado para cache-busting de CSS/JS sem exigir restart do servidor."""
    try:
        candidates = [
            BASE_DIR / "ui" / "static" / "app.js",
            BASE_DIR / "ui" / "static" / "style.css",
        ]
        mtimes = [path.stat().st_mtime for path in candidates if path.exists()]
        return str(int(max(mtimes))) if mtimes else "1"
    except Exception:
        return "1"


@app.middleware("http")
async def token_auth_middleware(request: Request, call_next):
    if not API_TOKEN or _is_public_path(request.url.path):
        return await call_next(request)
    provided_token = _extract_token(request)
    if provided_token and secrets.compare_digest(provided_token, API_TOKEN):
        return await call_next(request)
    if request.url.path.startswith("/api/"):
        return JSONResponse({"error": "Nao autorizado."}, status_code=401)
    return HTMLResponse("Acesso nao autorizado.", status_code=401)


def _terminate_process(proc: subprocess.Popen[str]) -> None:
    if proc.poll() is not None:
        return
    try:
        proc.terminate()
        proc.wait(timeout=2)
    except Exception:
        try:
            proc.kill()
        except Exception:
            pass


class JobStoppedError(RuntimeError):
    pass


def _run_job(job: Job) -> None:
    env = os.environ.copy()
    env["PYTHONPATH"] = str(BASE_DIR)
    env["ORIGEM_EXECUCAO"] = str(job.trigger or "MANUAL").strip() or "MANUAL"
    cmd = [sys.executable, "-m", "App.main", "--input", str(job.input_path), "--output", str(job.output_path)]
    if job.email_extra_attachment:
        extra_path = _normalize_email_attachment_path(job.email_extra_attachment)
        if extra_path:
            cmd.extend(["--input-ml", str(extra_path)])
    proc = subprocess.Popen(
        cmd,
        cwd=str(BASE_DIR),
        env=env,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    with _runner_lock:
        global _current_process, _current_running_job_id
        _current_process = proc
        _current_running_job_id = job.job_id

    try:
        stdout, stderr = proc.communicate()
    finally:
        with _runner_lock:
            _current_process = None
            _current_running_job_id = None

    if job.job_id in _stop_requested_jobs:
        _stop_requested_jobs.discard(job.job_id)
        raise JobStoppedError("Execucao interrompida manualmente.")

    if proc.returncode != 0:
        err = (stderr or stdout or "").strip()
        raise RuntimeError(f"App.main falhou (exit={proc.returncode}): {(err or 'sem detalhes')[-2000:]}")

    # Se houver planilha secundaria, consolidar e substituir o output principal.
    if job.email_extra_attachment:
        merged_path, _ = _build_merged_output(job)
        if merged_path and merged_path.exists():
            try:
                shutil.copy2(merged_path, job.output_path)
            except Exception:
                pass


def _send_output_email(job: Job, recipients_override: list[str] | None = None) -> tuple[str, str | None]:
    if not EMAIL_ENABLED:
        return "SKIPPED", "Envio de e-mail desativado (EMAIL_ENABLED=false)."
    recipients = recipients_override or list(job.email_recipients or [])
    if not recipients:
        return "SKIPPED", "Nenhum destinatario informado."
    if not job.output_path or not job.output_path.exists():
        return "FAILED", "Arquivo de output nao encontrado."
    extra_warning = None
    msg = EmailMessage()
    subject_prefix = EMAIL_SUBJECT_PREFIX.strip() or "[Price Monitor]"
    msg["Subject"] = f"{subject_prefix} output job {job.job_id}"
    msg["From"] = SMTP_SENDER
    msg["To"] = ", ".join(recipients)
    msg.set_content(f"Segue output do job {job.job_id} em anexo.")
    base_path = job.output_path
    primary_path = base_path
    if job.output_mode == "a_prazo":
        simple_path, simple_error = _build_simple_output(base_path, job.output_mode)
        if simple_path:
            primary_path = simple_path
        else:
            extra_warning = simple_error or "Falha ao gerar planilha a prazo; enviado output completo."
    msg.add_attachment(
        primary_path.read_bytes(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=primary_path.name,
    )
    extra_attachment = _normalize_email_attachment_path(job.email_extra_attachment)
    if job.email_extra_attachment and not extra_attachment and not extra_warning:
        extra_warning = "Anexo extra nao encontrado; enviado apenas o output."

    smtp_ready = bool(SMTP_HOST and SMTP_SENDER)

    def _send_via_outlook() -> tuple[str, str | None]:
        if os.name != "nt":
            return "FAILED", "Modo Outlook disponivel apenas no Windows."
        subject = msg["Subject"] or f"{EMAIL_SUBJECT_PREFIX} output job {job.job_id}"
        body = f"Segue output do job {job.job_id} em anexo."
        attachment_path = str(primary_path.resolve())
        extra_path = ""

        try:
            import win32com.client  # type: ignore
            outlook = win32com.client.Dispatch("Outlook.Application")
            mail = outlook.CreateItem(0)  # 0 = olMailItem
            mail.To = "; ".join(recipients)
            mail.Subject = subject
            mail.Body = body
            mail.Attachments.Add(attachment_path)
            mail.Send()
            return "SENT", extra_warning
        except Exception as exc:
            # Fallback sem pywin32: usa COM pelo PowerShell (nativo no Windows).
            env = os.environ.copy()
            env["PM_TO"] = "; ".join(recipients)
            env["PM_SUBJECT"] = subject
            env["PM_BODY"] = body
            env["PM_ATTACHMENT"] = attachment_path
            ps_script = (
                "$ErrorActionPreference='Stop';"
                "$outlook = New-Object -ComObject Outlook.Application;"
                "$mail = $outlook.CreateItem(0);"
                "$mail.To = $env:PM_TO;"
                "$mail.Subject = $env:PM_SUBJECT;"
                "$mail.Body = $env:PM_BODY;"
                "if ($env:PM_ATTACHMENT -and (Test-Path $env:PM_ATTACHMENT)) { $null = $mail.Attachments.Add($env:PM_ATTACHMENT) };"
                "$mail.Send()"
            )
            try:
                result = subprocess.run(
                    ["powershell", "-NoProfile", "-Command", ps_script],
                    env=env,
                    capture_output=True,
                    text=True,
                    timeout=45,
                )
            except Exception as ps_exc:
                return "FAILED", f"Falha Outlook/PowerShell: {type(ps_exc).__name__}: {ps_exc}"

            if result.returncode == 0:
                return "SENT", extra_warning

            stderr = (result.stderr or "").strip()
            stdout = (result.stdout or "").strip()
            details = stderr or stdout or "sem detalhes"
            first_error = f"{type(exc).__name__}: {exc}"
            return "FAILED", f"Falha Outlook COM ({first_error}); fallback PowerShell falhou: {details}"

    def _send_via_smtp() -> tuple[str, str | None]:
        if not SMTP_HOST or not SMTP_SENDER:
            hint = (
                "SMTP nao configurado. "
                "Para Gmail, defina no .env: GMAIL_USER e GMAIL_APP_PASSWORD "
                "(App Password do Google), e opcionalmente GMAIL_SENDER."
            )
            return "SKIPPED", hint

        last_exc: Exception | None = None
        for attempt in range(1, SMTP_RETRY_ATTEMPTS + 1):
            try:
                if SMTP_USE_SSL:
                    with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=30, context=ssl.create_default_context()) as server:
                        if SMTP_USER and SMTP_PASSWORD:
                            server.login(SMTP_USER, SMTP_PASSWORD)
                        server.send_message(msg)
                else:
                    with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=30) as server:
                        if SMTP_USE_TLS:
                            server.starttls(context=ssl.create_default_context())
                        if SMTP_USER and SMTP_PASSWORD:
                            server.login(SMTP_USER, SMTP_PASSWORD)
                        server.send_message(msg)
                return "SENT", extra_warning
            except Exception as exc:
                last_exc = exc
                if attempt < SMTP_RETRY_ATTEMPTS:
                    delay = min(SMTP_RETRY_MAX_SECONDS, SMTP_RETRY_MIN_SECONDS * (2 ** (attempt - 1)))
                    time.sleep(delay)

        if last_exc is None:
            return "FAILED", "Falha SMTP: erro desconhecido."

        hint = ""
        if isinstance(last_exc, OSError):
            winerror = getattr(last_exc, "winerror", None)
            if winerror == 10013:
                hint = " Acesso a socket bloqueado (firewall/antivirus/politica local)."
            elif winerror == 10061:
                hint = " Conexao recusada pelo servidor SMTP."
        return "FAILED", f"Falha SMTP apos {SMTP_RETRY_ATTEMPTS} tentativa(s): {type(last_exc).__name__}: {last_exc}.{hint}"

    if EMAIL_DELIVERY_MODE == "smtp":
        return _send_via_smtp()

    if EMAIL_DELIVERY_MODE == "outlook":
        return _send_via_outlook()

    # Modo automatico:
    # - Prioriza SMTP (inclui Gmail via SMTP).
    # - Nao tenta Outlook a menos que OUTLOOK_FALLBACK_ENABLED=true (ou modo outlook explicito).
    if not smtp_ready:
        if OUTLOOK_FALLBACK_ENABLED:
            outlook_status, outlook_error = _send_via_outlook()
            if outlook_status == "SENT":
                return outlook_status, None
            return "FAILED", f"SMTP: nao configurado | OUTLOOK: {outlook_error}"
        return _send_via_smtp()

    smtp_status, smtp_error = _send_via_smtp()
    if smtp_status == "SENT":
        return smtp_status, smtp_error

    if not OUTLOOK_FALLBACK_ENABLED:
        return "FAILED", smtp_error

    outlook_status, outlook_error = _send_via_outlook()
    if outlook_status == "SENT":
        return outlook_status, None

    combined = f"SMTP: {smtp_error} | OUTLOOK: {outlook_error}"
    return "FAILED", combined


def _build_public_download_url(job: Job, merged: bool = False, output_mode: str | None = None) -> str | None:
    base = (APP_PUBLIC_BASE_URL or "").strip().rstrip("/")
    if not base:
        return None

    url = f"{base}/download/{job.job_id}"
    if merged:
        sep = "&" if "?" in url else "?"
        url = f"{url}{sep}merged=1"
    if output_mode == "a_prazo":
        sep = "&" if "?" in url else "?"
        url = f"{url}{sep}mode=aprazo"
    if API_TOKEN:
        sep = "&" if "?" in url else "?"
        url = f"{url}{sep}{urlencode({'api_token': API_TOKEN})}"
    return url


def _send_output_whatsapp(job: Job, recipients_override: list[str] | None = None) -> tuple[str, str | None]:
    recipients = recipients_override or list(job.whatsapp_recipients or [])
    if not recipients:
        return "SKIPPED", "Nenhum numero WhatsApp informado."
    if not job.output_path or not job.output_path.exists():
        return "FAILED", "Arquivo de output nao encontrado."

    from_number = _normalize_whatsapp_recipient(TWILIO_WHATSAPP_FROM)
    if not (TWILIO_ACCOUNT_SID and TWILIO_AUTH_TOKEN and from_number):
        return "FAILED", "Twilio WhatsApp nao configurado (SID/TOKEN/FROM)."

    use_merged = bool(job.email_extra_attachment)
    public_url = _build_public_download_url(job, merged=use_merged, output_mode=job.output_mode)
    if public_url:
        body = f"{WHATSAPP_MESSAGE_PREFIX} Job {job.job_id} concluido. Planilha: {public_url}"
    else:
        body = (
            f"{WHATSAPP_MESSAGE_PREFIX} Job {job.job_id} concluido. "
            "Planilha disponivel no painel (configure APP_PUBLIC_BASE_URL para enviar link direto)."
        )

    endpoint = f"https://api.twilio.com/2010-04-01/Accounts/{TWILIO_ACCOUNT_SID}/Messages.json"
    errors: list[str] = []
    sent = 0
    for recipient in recipients:
        payload = {
            "From": from_number,
            "To": recipient,
            "Body": body,
        }
        if WHATSAPP_MEDIA_ENABLED and public_url:
            payload["MediaUrl"] = public_url
        try:
            response = requests.post(
                endpoint,
                data=payload,
                auth=(TWILIO_ACCOUNT_SID, TWILIO_AUTH_TOKEN),
                timeout=30,
            )
            if response.status_code >= 400:
                detail = response.text.strip()
                errors.append(f"{recipient}: HTTP {response.status_code} {detail[:220]}")
                continue
            sent += 1
        except Exception as exc:
            errors.append(f"{recipient}: {type(exc).__name__}: {exc}")

    if sent == 0:
        return "FAILED", "; ".join(errors) if errors else "Falha desconhecida no envio WhatsApp."
    if errors:
        return "FAILED", f"Parcial ({sent}/{len(recipients)}): {'; '.join(errors)}"
    return "SENT", None


def _send_output_drive(
    job: Job,
    folder_id_override: str | None = None,
    credentials_file_override: str | None = None,
) -> tuple[str, str | None, str | None, str | None]:
    if not job.output_path or not job.output_path.exists():
        return "FAILED", "Arquivo de output nao encontrado.", None, None
    merged_path = None
    if job.email_extra_attachment:
        merged_path, _ = _build_merged_output(job)

    folder_id = _normalize_drive_folder_id(folder_id_override or job.drive_folder_id)
    if not folder_id:
        return "FAILED", "Pasta do Google Drive nao configurada.", None, None

    credentials_file = _normalize_drive_credentials_file(credentials_file_override or job.drive_credentials_file)
    if not credentials_file:
        return "FAILED", "Arquivo de credenciais do Google Drive nao configurado.", None, None

    credentials_path = Path(credentials_file)
    if not credentials_path.exists():
        return "FAILED", f"Arquivo de credenciais nao encontrado: {credentials_file}", None, None

    try:
        from google.oauth2 import service_account  # type: ignore
        from googleapiclient.discovery import build  # type: ignore
        from googleapiclient.http import MediaFileUpload  # type: ignore
    except Exception as exc:
        return "FAILED", f"Dependencias do Google Drive indisponiveis: {type(exc).__name__}: {exc}", None, None

    try:
        credentials = service_account.Credentials.from_service_account_file(
            str(credentials_path),
            scopes=["https://www.googleapis.com/auth/drive"],
        )
        service = build("drive", "v3", credentials=credentials, cache_discovery=False)
        base_path = merged_path or job.output_path
        upload_path = str(base_path)
        if job.output_mode == "a_prazo":
            simple_path, simple_error = _build_simple_output(base_path, job.output_mode)
            if simple_path:
                upload_path = str(simple_path)
            else:
                return "FAILED", simple_error or "Falha ao gerar planilha a prazo.", None, None
        media = MediaFileUpload(
            upload_path,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            resumable=False,
        )
        filename = Path(upload_path).name
        body = {"name": filename, "parents": [folder_id]}
        created = service.files().create(
            body=body,
            media_body=media,
            fields="id, webViewLink, webContentLink",
        ).execute()
    except Exception as exc:
        return "FAILED", f"Falha no upload para Google Drive: {type(exc).__name__}: {exc}", None, None

    file_id = str(created.get("id") or "").strip() or None
    file_url = str(created.get("webViewLink") or created.get("webContentLink") or "").strip() or None
    if file_id and not file_url:
        file_url = f"https://drive.google.com/file/d/{file_id}/view"
    return "SENT", None, file_id, file_url


def _worker() -> None:
    while True:
        job_id = job_queue.get()
        job = jobs.get(job_id)
        if job is None:
            job_queue.task_done()
            continue

        job.status = "RUNNING"
        job.started_at = _now_ts()
        try:
            _run_job(job)
            job.status = "DONE"
            _archive_result_by_origin(job)
        except JobStoppedError as exc:
            job.status = "STOPPED"
            job.error = str(exc)
        except Exception as exc:
            job.status = "FAILED"
            job.error = f"{type(exc).__name__}: {exc}"
        finally:
            job.finished_at = _now_ts()
            if job.status == "DONE" and job.auto_email_enabled:
                job.email_status, job.email_error = _send_output_email(job)
                if job.email_status == "SENT":
                    job.email_sent_at = _now_ts()
            elif job.email_status == "PENDING":
                job.email_status = "SKIPPED"

            if job.status == "DONE" and job.auto_whatsapp_enabled:
                job.whatsapp_status, job.whatsapp_error = _send_output_whatsapp(job)
                if job.whatsapp_status == "SENT":
                    job.whatsapp_sent_at = _now_ts()
            elif job.whatsapp_status == "PENDING":
                job.whatsapp_status = "SKIPPED"

            if job.status == "DONE" and job.auto_drive_enabled:
                job.drive_status, job.drive_error, job.drive_file_id, job.drive_file_url = _send_output_drive(job)
                if job.drive_status == "SENT":
                    job.drive_uploaded_at = _now_ts()
            elif job.drive_status == "PENDING":
                job.drive_status = "SKIPPED"
            job_queue.task_done()


def _active_job_exists() -> bool:
    return any(job.status in {"QUEUED", "RUNNING"} for job in jobs.values())


def _enqueue_job(
    file: UploadFile | None,
    trigger: str,
    recipients: list[str] | None,
    auto_email_enabled: bool | None,
    email_extra_attachment: str | None = None,
    whatsapp_recipients: list[str] | None = None,
    auto_whatsapp_enabled: bool | None = None,
    drive_folder_id: str | None = None,
    drive_credentials_file: str | None = None,
    auto_drive_enabled: bool | None = None,
    output_mode: str | None = None,
    input_override_path: Path | None = None,
) -> tuple[Job | None, str | None]:
    global latest_manual_job_id, latest_daily_job_id
    if _active_job_exists():
        return None, ONLY_ONE_JOB_MESSAGE

    job_id = _new_job_id(trigger=trigger)
    job_dir = RUNS_DIR / job_id
    job_dir.mkdir(parents=True, exist_ok=True)
    input_path = job_dir / "input.xlsx"
    output_path = job_dir / _output_filename(job_id)

    if input_override_path is not None:
        if not input_override_path.exists():
            return None, f"Arquivo de entrada para processamento nao encontrado: {input_override_path}"
        shutil.copy(input_override_path, input_path)
    elif file is not None:
        with input_path.open("wb") as f:
            shutil.copyfileobj(file.file, f)
        try:
            file.file.close()
        except Exception:
            pass
    else:
        if trigger != "MANUAL":
            if PRIMARY_SCHEDULED_INPUT.exists():
                shutil.copy(PRIMARY_SCHEDULED_INPUT, input_path)
            elif DEFAULT_INPUT.exists():
                shutil.copy(DEFAULT_INPUT, input_path)
            else:
                return None, "Planilha principal nao encontrada para execucao automatica."
        else:
            if not DEFAULT_INPUT.exists():
                return None, "Nenhuma planilha enviada e input.xlsx nao encontrado."
            shutil.copy(DEFAULT_INPUT, input_path)

    if not email_extra_attachment:
        latest_secondary = _resolve_secondary_attachment()
        if latest_secondary is not None:
            email_extra_attachment = str(latest_secondary)
        elif SECONDARY_ML_REQUIRED:
            return None, f"Planilha secundaria nao encontrada em {SECONDARY_INPUT_DIR}."

    if email_extra_attachment:
        extra_src = Path(_normalize_email_attachment_path(email_extra_attachment))
        if extra_src.exists():
            extra_target = job_dir / extra_src.name
            try:
                shutil.copy(extra_src, extra_target)
                email_extra_attachment = str(extra_target)
            except Exception:
                pass

    with _email_settings_lock:
        base_recipients = list(recipients) if recipients is not None else list(_runtime_email_recipients)
        effective_recipients = _merge_recipients(base_recipients, ALWAYS_EMAIL_RECIPIENTS)
        # sempre enviar ao finalizar
        if auto_email_enabled is None:
            effective_auto_email = True
        else:
            effective_auto_email = bool(auto_email_enabled) or True
        effective_extra_attachment = _normalize_email_attachment_path(email_extra_attachment)

    with _whatsapp_settings_lock:
        effective_whatsapp_recipients = (
            list(whatsapp_recipients) if whatsapp_recipients is not None else list(_runtime_whatsapp_recipients)
        )
        if auto_whatsapp_enabled is None:
            effective_auto_whatsapp = _runtime_auto_whatsapp_manual if trigger == "MANUAL" else _runtime_auto_whatsapp_daily
        else:
            effective_auto_whatsapp = bool(auto_whatsapp_enabled)

    with _drive_settings_lock:
        effective_drive_folder_id = _normalize_drive_folder_id(
            drive_folder_id if drive_folder_id is not None else _runtime_drive_folder_id
        )
        effective_drive_credentials_file = _normalize_drive_credentials_file(
            drive_credentials_file if drive_credentials_file is not None else _runtime_drive_credentials_file
        )
        if auto_drive_enabled is None:
            effective_auto_drive = _runtime_auto_drive_manual if trigger == "MANUAL" else _runtime_auto_drive_daily
        else:
            effective_auto_drive = bool(auto_drive_enabled)

    normalized_output_mode = str(output_mode or "").strip().lower()
    if normalized_output_mode not in {"completa", "a_prazo", "a_vista"}:
        normalized_output_mode = "a_prazo"
    if normalized_output_mode == "a_vista" and not FEATURE_AVISTA_ENABLED:
        normalized_output_mode = "a_prazo"

    job = Job(
        job_id=job_id,
        trigger=trigger,
        input_path=input_path,
        output_path=output_path,
        email_recipients=effective_recipients,
        auto_email_enabled=effective_auto_email,
        email_extra_attachment=effective_extra_attachment or "",
        whatsapp_recipients=effective_whatsapp_recipients,
        auto_whatsapp_enabled=effective_auto_whatsapp,
        drive_folder_id=effective_drive_folder_id,
        drive_credentials_file=effective_drive_credentials_file,
        auto_drive_enabled=effective_auto_drive,
        output_mode=normalized_output_mode,
    )
    jobs[job_id] = job
    job_queue.put(job_id)

    if trigger == "MANUAL":
        latest_manual_job_id = job_id
    else:
        latest_daily_job_id = job_id
    return job, None


def _next_daily_run_dt() -> datetime:
    run_times = _daily_run_times()
    now = datetime.now()
    today = now.date()
    for hour, minute in run_times:
        candidate = datetime(today.year, today.month, today.day, hour, minute)
        if candidate > now:
            return candidate
    h, m = run_times[0]
    tomorrow = now + timedelta(days=1)
    return datetime(tomorrow.year, tomorrow.month, tomorrow.day, h, m)


def _daily_run_times() -> list[tuple[int, int]]:
    out = [_fixed_daily_run_slot()]
    for slot in _extra_daily_run_times():
        if _daily_slot_enabled(_format_daily_run_time(slot)):
            out.append(slot)
    return sorted(set(out))


def _daily_run_times_text() -> list[str]:
    return [_format_daily_run_time(value) for value in _daily_run_times()]


def _daily_auto_enabled() -> bool:
    return bool(_parse_bool(os.getenv("DAILY_AUTO_ENABLED", "1"), default=True))


def _daily_trigger_window_seconds() -> int:
    raw = os.getenv("DAILY_TRIGGER_WINDOW_SECONDS", "90")
    try:
        value = int(str(raw).strip())
    except Exception:
        value = 90
    if value < 10:
        return 10
    if value > 900:
        return 900
    return value


def _current_daily_slot(now: datetime, window_seconds: int) -> tuple[int, int] | None:
    for hour, minute in _daily_run_times():
        target = datetime(now.year, now.month, now.day, hour, minute)
        delta = (now - target).total_seconds()
        if 0 <= delta <= window_seconds:
            return hour, minute
    return None


def _daily_scheduler_loop() -> None:
    while True:
        try:
            if not _daily_auto_enabled():
                time.sleep(5)
                continue

            now = datetime.now()
            slot = _current_daily_slot(now, _daily_trigger_window_seconds())
            if not slot:
                time.sleep(5)
                continue

            slot_key = f"{now:%Y%m%d}-{slot[0]:02d}:{slot[1]:02d}"
            with _daily_scheduler_lock:
                today_prefix = now.strftime("%Y%m%d-")
                _daily_triggered_slots.intersection_update({k for k in _daily_triggered_slots if k.startswith(today_prefix)})
                already_triggered = slot_key in _daily_triggered_slots

            if already_triggered:
                time.sleep(5)
                continue

            job, error = _enqueue_job(
                file=None,
                trigger="AUTO_DIARIO",
                recipients=None,
                auto_email_enabled=None,
                email_extra_attachment=None,
            )
            if job is not None:
                with _daily_scheduler_lock:
                    _daily_triggered_slots.add(slot_key)
            elif error != ONLY_ONE_JOB_MESSAGE:
                with _daily_scheduler_lock:
                    _daily_triggered_slots.add(slot_key)

            time.sleep(5)
        except Exception:
            time.sleep(5)


def _magalu_auth_url(request: Request, creds: dict[str, str]) -> str:
    client_id = str(creds.get("api_key_id") or "").strip()
    if not client_id:
        return ""
    redirect_uri = str(creds.get("redirect_uri") or "").strip() or f"{str(request.base_url).rstrip('/')}/api/marketplace/callback/magalu"
    state = str(creds.get("oauth_state") or "").strip() or secrets.token_urlsafe(24)
    if not creds.get("oauth_state"):
        creds["oauth_state"] = state
        _save_marketplace_credentials()
    params = {
        "response_type": "code",
        "client_id": client_id,
        "redirect_uri": redirect_uri,
        "scope": str(creds.get("auth_scope") or "openid profile email"),
        "state": state,
    }
    return f"{os.getenv('MAGALU_AUTHORIZATION_URL', 'https://id.magalu.com/oauth/authorize')}?{urlencode(params)}"


def _magalu_payload(request: Request) -> dict[str, Any]:
    with _marketplace_lock:
        creds = dict(_marketplace_credentials.get("magalu") or _default_magalu_credentials())
    return {
        "has_api_key": bool(creds.get("api_key")),
        "has_api_key_id": bool(creds.get("api_key_id")),
        "has_api_key_secret": bool(creds.get("api_key_secret")),
        "has_access_token": bool(creds.get("access_token")),
        "has_refresh_token": bool(creds.get("refresh_token")),
        "api_base": creds.get("api_base", ""),
        "token_url": creds.get("token_url", ""),
        "prices_path": creds.get("prices_path", ""),
        "seller_id": creds.get("seller_id", ""),
        "redirect_uri": creds.get("redirect_uri", ""),
        "return_url": str(creds.get("redirect_uri") or "").strip() or f"{str(request.base_url).rstrip('/')}/api/marketplace/callback/magalu",
        "auth_scope": creds.get("auth_scope", ""),
        "authorization_url": _magalu_auth_url(request, creds),
    }


def read_output(path: Path):
    if not path.exists():
        return [], []
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    rows = [row for row in ws.iter_rows(min_row=2, values_only=True)]
    wb.close()
    return headers, rows


def _append_precos_import_log(entry: dict[str, Any]) -> None:
    RUNS_DIR.mkdir(exist_ok=True)
    line = json.dumps(entry, ensure_ascii=False)
    with _precos_import_log_lock:
        with PRECOS_IMPORT_LOG_FILE.open("a", encoding="utf-8") as fh:
            fh.write(line + "\n")


def _new_ml_task_id() -> str:
    while True:
        task_id = f"mlt_{secrets.token_hex(6)}"
        if task_id not in _ml_tasks:
            return task_id


def _save_ml_tasks() -> None:
    _write_json(ML_TASKS_FILE, {"tasks": _ml_tasks})


def _load_ml_tasks() -> None:
    global _ml_tasks
    data = _read_json(ML_TASKS_FILE, {"tasks": {}})
    tasks = data.get("tasks")
    if not isinstance(tasks, dict):
        _ml_tasks = {}
        return
    sanitized: dict[str, dict[str, Any]] = {}
    for task_id, raw in tasks.items():
        if not isinstance(task_id, str) or not isinstance(raw, dict):
            continue
        sanitized[task_id] = {
            "task_id": str(raw.get("task_id") or task_id),
            "status": str(raw.get("status") or "pending"),
            "entries": [str(x).strip() for x in (raw.get("entries") or []) if str(x).strip()],
            "limite_por_pesquisa": int(raw.get("limite_por_pesquisa") or 10),
            "delay_ms": int(raw.get("delay_ms") or 1200),
            "created_at": float(raw.get("created_at") or _now_ts()),
            "started_at": raw.get("started_at"),
            "finished_at": raw.get("finished_at"),
            "attempt_count": int(raw.get("attempt_count") or 0),
            "result_count": int(raw.get("result_count") or 0),
            "erro": str(raw.get("erro") or ""),
        }
    _ml_tasks = sanitized


def _next_pending_ml_task() -> dict[str, Any] | None:
    pending = [task for task in _ml_tasks.values() if task.get("status") == "pending"]
    if not pending:
        return None
    pending.sort(key=lambda item: float(item.get("created_at") or 0))
    return pending[0]


def _create_ml_task(entries: list[str], limite_por_pesquisa: int = 10, delay_ms: int = 1200) -> dict[str, Any]:
    cleaned = [str(item or "").strip() for item in entries if str(item or "").strip()]
    if not cleaned:
        raise ValueError("Nenhuma entrada valida para tarefa de Mercado Livre.")

    with _ml_tasks_lock:
        task_id = _new_ml_task_id()
        task = {
            "task_id": task_id,
            "status": "pending",
            "entries": cleaned,
            "limite_por_pesquisa": max(1, min(50, int(limite_por_pesquisa or 10))),
            "delay_ms": max(300, int(delay_ms or 1200)),
            "created_at": _now_ts(),
            "started_at": None,
            "finished_at": None,
            "attempt_count": 0,
            "result_count": 0,
            "erro": "",
        }
        _ml_tasks[task_id] = task
        _save_ml_tasks()
        return task


def _extract_ml_entries_from_input(path: Path) -> list[str]:
    if not path.exists():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [texto if isinstance(texto := cell, str) else str(cell or "") for cell in rows[0]]
        normalized_headers = []
        for h in headers:
            base = unicodedata.normalize("NFKD", str(h or ""))
            base = "".join(ch for ch in base if not unicodedata.combining(ch))
            normalized_headers.append(base.strip().lower())

        preferred_keys = (
            "link",
            "url",
            "pesquisa",
            "termo",
            "produto",
            "titulo",
        )
        preferred_cols: list[int] = []
        for idx, header in enumerate(normalized_headers):
            if any(key in header for key in preferred_keys):
                preferred_cols.append(idx)
        if not preferred_cols:
            preferred_cols = [0]

        out: list[str] = []
        seen: set[str] = set()
        for row in rows[1:]:
            for col_idx in preferred_cols:
                if col_idx >= len(row):
                    continue
                raw = row[col_idx]
                text = str(raw or "").strip()
                if not text:
                    continue
                key = text.lower()
                if key in seen:
                    continue
                seen.add(key)
                out.append(text)
        return out
    finally:
        wb.close()


def _looks_like_mercadolivre(value: object) -> bool:
    text = str(value or "").strip().lower()
    return "mercadolivre.com.br" in text or "produto.mercadolivre.com.br" in text


def _extract_ml_entries_and_strip_rows(path: Path) -> tuple[list[str], int, int]:
    if not path.exists():
        return [], 0, 0

    wb = load_workbook(path)
    try:
        ws = wb.active
        if ws.max_row < 2:
            return [], 0, 0

        header_map: dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            name = str(ws.cell(row=1, column=col).value or "").strip().lower()
            if name:
                header_map[name] = col

        link_col = None
        for candidate in ("link", "url", "href"):
            if candidate in header_map:
                link_col = header_map[candidate]
                break
        canal_col = header_map.get("canal")

        ml_entries: list[str] = []
        rows_to_delete: list[int] = []
        seen: set[str] = set()

        for row in range(2, ws.max_row + 1):
            link_value = str(ws.cell(row=row, column=link_col).value or "").strip() if link_col else ""
            canal_value = str(ws.cell(row=row, column=canal_col).value or "").strip() if canal_col else ""

            is_ml = _looks_like_mercadolivre(link_value) or "mercado livre" in canal_value.lower()
            if not is_ml:
                continue

            entry = link_value or canal_value
            entry_key = entry.lower().strip()
            if entry_key and entry_key not in seen:
                seen.add(entry_key)
                ml_entries.append(entry)

            rows_to_delete.append(row)

        for row in reversed(rows_to_delete):
            ws.delete_rows(row, 1)

        if rows_to_delete:
            wb.save(path)

        remaining_data_rows = max(0, ws.max_row - 1)
        return ml_entries, len(rows_to_delete), remaining_data_rows
    finally:
        wb.close()


@app.on_event("startup")
def startup() -> None:
    RUNS_DIR.mkdir(exist_ok=True)
    _load_email_settings()
    _load_whatsapp_settings()
    _load_drive_settings()
    _load_daily_schedule_settings()
    _load_users()
    _load_marketplace_credentials()
    _load_ml_tasks()
    Thread(target=_worker, daemon=True).start()
    Thread(target=_daily_scheduler_loop, daemon=True).start()

@app.get("/api/health")
def api_health():
    return JSONResponse({"status": "ok"})


@app.get("/api/overview")
def api_overview():
    running = None
    with _runner_lock:
        if _current_running_job_id:
            running = jobs.get(_current_running_job_id)
    return JSONResponse(
        {
            "counts": _counts(),
            "has_default_input": DEFAULT_INPUT.exists(),
            "current_job": _job_payload(running),
            "latest_manual_job": _job_payload(_latest_manual_job()),
            "latest_daily_job": _job_payload(_latest_daily_job()),
            "latest_output_job": _job_payload(_latest_output_job()),
        }
    )


@app.post("/api/precos/importar")
def api_precos_importar(payload: PrecosImportPayload, request: Request):
    if not payload.itens:
        return JSONResponse({"error": "Payload sem itens para importar."}, status_code=400)

    invalid_indexes: list[int] = []
    for idx, item in enumerate(payload.itens):
        if item.preco <= 0:
            invalid_indexes.append(idx)

    if invalid_indexes:
        return JSONResponse(
            {
                "error": "Foram encontrados precos invalidos (<= 0).",
                "invalid_item_indexes": invalid_indexes,
            },
            status_code=400,
        )

    price_count = len(payload.itens)
    price_min = min(item.preco for item in payload.itens)
    price_max = max(item.preco for item in payload.itens)
    price_avg = sum(item.preco for item in payload.itens) / float(price_count)
    now_utc = datetime.utcnow().isoformat(timespec="seconds") + "Z"

    try:
        _append_precos_import_log(
            {
                "received_at_utc": now_utc,
                "client_ip": request.client.host if request.client else "",
                "fonte": payload.fonte,
                "coletado_em": payload.coletado_em,
                "itens_recebidos": price_count,
                "resumo_precos": {
                    "min": round(price_min, 2),
                    "max": round(price_max, 2),
                    "avg": round(price_avg, 2),
                },
                "sample": payload.itens[0].model_dump(),
            }
        )
    except Exception:
        # Falha de auditoria nao deve interromper a importacao.
        pass

    return JSONResponse(
        {
            "status": "ok",
            "message": "Importacao recebida com sucesso.",
            "fonte": payload.fonte,
            "coletado_em": payload.coletado_em,
            "itens_recebidos": price_count,
            "resumo_precos": {
                "min": round(price_min, 2),
                "max": round(price_max, 2),
                "avg": round(price_avg, 2),
            },
            "sample": payload.itens[0].model_dump(),
        }
    )


@app.post("/api/tarefas/ml/criar")
def api_ml_tarefas_criar(payload: MLTarefaCreatePayload):
    entries = [str(item or "").strip() for item in payload.entries if str(item or "").strip()]
    if not entries:
        return JSONResponse({"error": "Informe ao menos 1 URL/termo em entries."}, status_code=400)
    task = _create_ml_task(entries, limite_por_pesquisa=payload.limite_por_pesquisa, delay_ms=payload.delay_ms)

    return JSONResponse({"status": "created", "task": task})


@app.get("/api/tarefas/ml/proxima")
def api_ml_tarefas_proxima():
    with _ml_tasks_lock:
        task = _next_pending_ml_task()
        if not task:
            return JSONResponse({"status": "empty", "message": "Sem tarefas pendentes no momento."})

        task["status"] = "in_progress"
        task["started_at"] = _now_ts()
        task["attempt_count"] = int(task.get("attempt_count") or 0) + 1
        _save_ml_tasks()
        out = dict(task)

    return JSONResponse({"status": "ok", "task": out})


@app.post("/api/tarefas/ml/{task_id}/concluir")
def api_ml_tarefas_concluir(task_id: str, payload: MLTarefaResultadoPayload):
    with _ml_tasks_lock:
        task = _ml_tasks.get(task_id)
        if not task:
            return JSONResponse({"error": "Tarefa nao encontrada."}, status_code=404)

        task["status"] = "done"
        task["finished_at"] = _now_ts()
        task["result_count"] = len(payload.itens or [])
        task["erro"] = str(payload.erro or "")
        _save_ml_tasks()
        out = dict(task)

    return JSONResponse({"status": "ok", "task": out})


@app.post("/api/tarefas/ml/{task_id}/falhar")
def api_ml_tarefas_falhar(task_id: str, payload: MLTarefaResultadoPayload):
    with _ml_tasks_lock:
        task = _ml_tasks.get(task_id)
        if not task:
            return JSONResponse({"error": "Tarefa nao encontrada."}, status_code=404)

        task["status"] = "failed"
        task["finished_at"] = _now_ts()
        task["result_count"] = len(payload.itens or [])
        task["erro"] = str(payload.erro or "Falha informada pelo coletor.")
        _save_ml_tasks()
        out = dict(task)

    return JSONResponse({"status": "ok", "task": out})


@app.get("/api/email/settings")
def api_email_settings_get():
    return JSONResponse(_public_email_settings())


@app.post("/api/email/settings")
def api_email_settings_post(payload: EmailSettingsPayload):
    recipients = _parse_recipients(payload.recipients)
    if payload.recipients.strip() and not recipients:
        return JSONResponse({"error": "Nenhum e-mail valido informado."}, status_code=400)
    with _email_settings_lock:
        global _runtime_email_recipients, _runtime_auto_email_manual, _runtime_auto_email_daily
        _runtime_email_recipients = recipients
        _runtime_auto_email_manual = bool(payload.auto_send_manual)
        _runtime_auto_email_daily = bool(payload.auto_send_daily)
    return JSONResponse(_public_email_settings())


@app.get("/api/whatsapp/settings")
def api_whatsapp_settings_get():
    return JSONResponse(_public_whatsapp_settings())


@app.post("/api/whatsapp/settings")
def api_whatsapp_settings_post(payload: WhatsAppSettingsPayload):
    recipients = _parse_whatsapp_recipients(payload.recipients)
    if payload.recipients.strip() and not recipients:
        return JSONResponse({"error": "Nenhum numero WhatsApp valido informado."}, status_code=400)
    with _whatsapp_settings_lock:
        global _runtime_whatsapp_recipients, _runtime_auto_whatsapp_manual, _runtime_auto_whatsapp_daily
        _runtime_whatsapp_recipients = recipients
        _runtime_auto_whatsapp_manual = bool(payload.auto_send_manual)
        _runtime_auto_whatsapp_daily = bool(payload.auto_send_daily)
    return JSONResponse(_public_whatsapp_settings())


@app.get("/api/drive/settings")
def api_drive_settings_get():
    return JSONResponse(_public_drive_settings())


@app.post("/api/drive/settings")
def api_drive_settings_post(payload: DriveSettingsPayload):
    folder_id = ""
    if payload.folder_id.strip():
        folder_id = _normalize_drive_folder_id(payload.folder_id)
        if not folder_id:
            return JSONResponse({"error": "Folder ID do Google Drive invalido."}, status_code=400)

    credentials_file = _normalize_drive_credentials_file(payload.credentials_file)
    with _drive_settings_lock:
        global _runtime_drive_credentials_file, _runtime_drive_folder_id, _runtime_auto_drive_manual, _runtime_auto_drive_daily
        _runtime_drive_credentials_file = credentials_file
        _runtime_drive_folder_id = folder_id
        _runtime_auto_drive_manual = bool(payload.auto_send_manual)
        _runtime_auto_drive_daily = bool(payload.auto_send_daily)
    _save_drive_settings()
    return JSONResponse(_public_drive_settings())


@app.post("/run")
def run_agent(file: UploadFile | None = File(default=None)):
    if ML_AUTO_ROUTE_ENABLED:
        try:
            source_path: Path | None = None
            if file is not None:
                route_dir = RUNS_DIR / "ml_task_inputs"
                route_dir.mkdir(parents=True, exist_ok=True)
                safe_name = Path(file.filename or "input.xlsx").name
                source_path = route_dir / f"{int(time.time())}_autoroute_form_{safe_name}"
                with source_path.open("wb") as f:
                    shutil.copyfileobj(file.file, f)
                try:
                    file.file.close()
                except Exception:
                    pass
                file = None
            elif DEFAULT_INPUT.exists():
                source_path = DEFAULT_INPUT

            if source_path is not None and source_path.exists():
                if source_path == DEFAULT_INPUT:
                    route_dir = RUNS_DIR / "ml_task_inputs"
                    route_dir.mkdir(parents=True, exist_ok=True)
                    copied = route_dir / f"{int(time.time())}_autoroute_form_input.xlsx"
                    shutil.copy(source_path, copied)
                    source_path = copied

                entries, _ml_removed, ml_remaining = _extract_ml_entries_and_strip_rows(source_path)
                if entries:
                    task = _create_ml_task(entries, limite_por_pesquisa=10, delay_ms=1200)
                    if ml_remaining == 0:
                        return HTMLResponse(
                            f"Tarefa Mercado Livre criada: {task['task_id']}. "
                            f"Aguardando processamento no Chrome/Tampermonkey.",
                            status_code=200,
                        )
                    file = None
                    job, error = _enqueue_job(
                        file=None,
                        trigger="MANUAL",
                        recipients=None,
                        auto_email_enabled=None,
                        email_extra_attachment=None,
                        input_override_path=source_path,
                    )
                    if job is None:
                        status = 429 if error == ONLY_ONE_JOB_MESSAGE else 400
                        return HTMLResponse(error or "Falha ao criar job.", status_code=status)
                    return RedirectResponse(f"/status/{job.job_id}", status_code=303)
        except Exception as exc:
            return HTMLResponse(f"Falha no roteamento automatico Mercado Livre: {type(exc).__name__}: {exc}", status_code=400)

    job, error = _enqueue_job(
        file=file,
        trigger="MANUAL",
        recipients=None,
        auto_email_enabled=None,
        email_extra_attachment=None,
    )
    if job is None:
        status = 429 if error == ONLY_ONE_JOB_MESSAGE else 400
        return HTMLResponse(error or "Falha ao criar job.", status_code=status)
    return RedirectResponse(f"/status/{job.job_id}", status_code=303)


@app.post("/api/run")
def api_run(
    file: UploadFile | None = File(default=None),
    email_recipients: str = Form(default=""),
    output_mode: str = Form(default="completa"),
    email_attachment_file: UploadFile | None = File(default=None),
    auto_email: str | None = Form(default=None),
    auto_email_daily: str | None = Form(default=None),
    whatsapp_recipients: str = Form(default=""),
    auto_whatsapp: str | None = Form(default=None),
    auto_whatsapp_daily: str | None = Form(default=None),
    drive_credentials_file: str = Form(default=""),
    drive_folder_id: str = Form(default=""),
    auto_drive: str | None = Form(default=None),
    auto_drive_daily: str | None = Form(default=None),
):
    requested_mode = str(output_mode or "completa").strip().lower()
    input_override_path: Path | None = None
    ml_task_auto: dict[str, Any] | None = None
    ml_rows_removed = 0
    ml_remaining_rows = 0

    if requested_mode == "tampermonkey_queue":
        temp_input_path: Path | None = None
        try:
            if file is not None:
                queue_dir = RUNS_DIR / "ml_task_inputs"
                queue_dir.mkdir(parents=True, exist_ok=True)
                safe_name = Path(file.filename or "input.xlsx").name
                temp_input_path = queue_dir / f"{int(time.time())}_{safe_name}"
                with temp_input_path.open("wb") as f:
                    shutil.copyfileobj(file.file, f)
                try:
                    file.file.close()
                except Exception:
                    pass
            else:
                if DEFAULT_INPUT.exists():
                    temp_input_path = DEFAULT_INPUT
                else:
                    return JSONResponse(
                        {"error": "Nenhuma planilha enviada e input.xlsx nao encontrado para criar fila Tampermonkey."},
                        status_code=400,
                    )

            entries = _extract_ml_entries_from_input(temp_input_path)
            if not entries:
                return JSONResponse(
                    {"error": "Nenhuma URL/termo encontrado na planilha para criar tarefa Tampermonkey."},
                    status_code=400,
                )
            task = _create_ml_task(entries, limite_por_pesquisa=10, delay_ms=1200)

            return JSONResponse(
                {
                    "mode": "tampermonkey_queue",
                    "status": "created",
                    "task_id": task["task_id"],
                    "entries_count": len(entries),
                    "task": task,
                }
            )
        except Exception as exc:
            return JSONResponse(
                {"error": f"Falha ao criar tarefa Tampermonkey: {type(exc).__name__}: {exc}"},
                status_code=400,
            )

    if ML_AUTO_ROUTE_ENABLED:
        try:
            source_path: Path | None = None
            if file is not None:
                route_dir = RUNS_DIR / "ml_task_inputs"
                route_dir.mkdir(parents=True, exist_ok=True)
                safe_name = Path(file.filename or "input.xlsx").name
                source_path = route_dir / f"{int(time.time())}_autoroute_{safe_name}"
                with source_path.open("wb") as f:
                    shutil.copyfileobj(file.file, f)
                try:
                    file.file.close()
                except Exception:
                    pass
                file = None
            elif DEFAULT_INPUT.exists():
                source_path = DEFAULT_INPUT

            if source_path is not None and source_path.exists():
                if source_path == DEFAULT_INPUT:
                    route_dir = RUNS_DIR / "ml_task_inputs"
                    route_dir.mkdir(parents=True, exist_ok=True)
                    copied = route_dir / f"{int(time.time())}_autoroute_input.xlsx"
                    shutil.copy(source_path, copied)
                    source_path = copied

                entries, ml_rows_removed, ml_remaining_rows = _extract_ml_entries_and_strip_rows(source_path)
                input_override_path = source_path
                if entries:
                    ml_task_auto = _create_ml_task(entries, limite_por_pesquisa=10, delay_ms=1200)
                    if ml_remaining_rows == 0:
                        return JSONResponse(
                            {
                                "mode": "tampermonkey_queue",
                                "status": "created",
                                "task_id": ml_task_auto["task_id"],
                                "entries_count": len(ml_task_auto.get("entries") or []),
                                "task": ml_task_auto,
                                "message": "Somente Mercado Livre detectado: tarefa enviada para Tampermonkey.",
                                "ml_auto_routed": True,
                                "ml_rows_removed_from_job": int(ml_rows_removed or 0),
                            }
                        )
        except Exception as exc:
            return JSONResponse(
                {"error": f"Falha no roteamento automatico do Mercado Livre: {type(exc).__name__}: {exc}"},
                status_code=400,
            )

    recipients = None
    if email_recipients.strip():
        recipients = _parse_recipients(email_recipients)
        if not recipients:
            return JSONResponse({"error": "Nenhum e-mail valido informado em email_recipients."}, status_code=400)
    extra_attachment = None
    if email_attachment_file is not None:
        try:
            filename = email_attachment_file.filename or ""
            if filename and not filename.lower().endswith(".xlsx"):
                return JSONResponse({"error": "Anexo extra deve ser .xlsx."}, status_code=400)
            attach_dir = RUNS_DIR / "extra_attachments"
            attach_dir.mkdir(parents=True, exist_ok=True)
            safe_name = Path(filename or "anexo.xlsx").name
            temp_path = attach_dir / f"{int(time.time())}_{safe_name}"
            with temp_path.open("wb") as handle:
                shutil.copyfileobj(email_attachment_file.file, handle)
            extra_attachment = str(temp_path)
            try:
                email_attachment_file.file.close()
            except Exception:
                pass
        except Exception:
            return JSONResponse({"error": "Falha ao salvar anexo extra enviado."}, status_code=400)
    if extra_attachment is None:
        latest_secondary = _resolve_secondary_attachment()
        if latest_secondary is not None:
            extra_attachment = str(latest_secondary)
        elif SECONDARY_ML_REQUIRED:
            return JSONResponse(
                {"error": f"Planilha secundária nao encontrada em {SECONDARY_INPUT_DIR}."},
                status_code=400,
            )

    wa_recipients = None
    if whatsapp_recipients.strip():
        wa_recipients = _parse_whatsapp_recipients(whatsapp_recipients)
        if not wa_recipients:
            return JSONResponse({"error": "Nenhum numero WhatsApp valido informado em whatsapp_recipients."}, status_code=400)

    auto_email_enabled = _parse_bool(auto_email, default=None)
    auto_daily_enabled = _parse_bool(auto_email_daily, default=None)
    auto_whatsapp_enabled = _parse_bool(auto_whatsapp, default=None)
    auto_whatsapp_daily_enabled = _parse_bool(auto_whatsapp_daily, default=None)
    auto_drive_enabled = _parse_bool(auto_drive, default=None)
    auto_drive_daily_enabled = _parse_bool(auto_drive_daily, default=None)

    effective_drive_folder_id = None
    if drive_folder_id.strip():
        effective_drive_folder_id = _normalize_drive_folder_id(drive_folder_id)
        if not effective_drive_folder_id:
            return JSONResponse({"error": "Folder ID do Google Drive invalido em drive_folder_id."}, status_code=400)

    effective_drive_credentials_file = None
    if drive_credentials_file.strip():
        effective_drive_credentials_file = _normalize_drive_credentials_file(drive_credentials_file)

    with _email_settings_lock:
        global _runtime_auto_email_daily
        if recipients is not None:
            global _runtime_email_recipients
            _runtime_email_recipients = list(recipients)
        if auto_daily_enabled is not None:
            _runtime_auto_email_daily = bool(auto_daily_enabled)

    with _whatsapp_settings_lock:
        global _runtime_auto_whatsapp_daily
        if wa_recipients is not None:
            global _runtime_whatsapp_recipients
            _runtime_whatsapp_recipients = list(wa_recipients)
        if auto_whatsapp_daily_enabled is not None:
            _runtime_auto_whatsapp_daily = bool(auto_whatsapp_daily_enabled)

    with _drive_settings_lock:
        global _runtime_drive_credentials_file, _runtime_drive_folder_id, _runtime_auto_drive_daily
        if effective_drive_credentials_file is not None:
            _runtime_drive_credentials_file = effective_drive_credentials_file
        if effective_drive_folder_id is not None:
            _runtime_drive_folder_id = effective_drive_folder_id
        if auto_drive_daily_enabled is not None:
            _runtime_auto_drive_daily = bool(auto_drive_daily_enabled)
    if (
        effective_drive_credentials_file is not None
        or effective_drive_folder_id is not None
        or auto_drive_daily_enabled is not None
    ):
        _save_drive_settings()

    job, error = _enqueue_job(
        file,
        "MANUAL",
        recipients,
        auto_email_enabled,
        email_extra_attachment=extra_attachment,
        whatsapp_recipients=wa_recipients,
        auto_whatsapp_enabled=auto_whatsapp_enabled,
        drive_folder_id=effective_drive_folder_id,
        drive_credentials_file=effective_drive_credentials_file,
        auto_drive_enabled=auto_drive_enabled,
        output_mode=output_mode,
        input_override_path=input_override_path,
    )
    if job is None:
        status = 429 if error == ONLY_ONE_JOB_MESSAGE else 400
        return JSONResponse({"error": error or "Falha ao criar job."}, status_code=status)

    return JSONResponse(
        {
            "job_id": job.job_id,
            "status_url": f"/status/{job.job_id}",
            "download_url": f"/download/{job.job_id}",
            "auto_email_enabled": job.auto_email_enabled,
            "email_recipients_csv": _recipients_to_csv(job.email_recipients),
            "email_extra_attachment": job.email_extra_attachment,
            "auto_whatsapp_enabled": job.auto_whatsapp_enabled,
            "whatsapp_recipients_csv": _recipients_to_csv(job.whatsapp_recipients),
            "auto_drive_enabled": job.auto_drive_enabled,
            "drive_folder_id": job.drive_folder_id,
            "ml_auto_routed": bool(ml_task_auto),
            "ml_task_id": str((ml_task_auto or {}).get("task_id") or ""),
            "ml_entries_count": len((ml_task_auto or {}).get("entries") or []),
            "ml_rows_removed_from_job": int(ml_rows_removed or 0),
        }
    )


@app.post("/api/jobs/stop")
def api_stop_running_job():
    with _runner_lock:
        running_id = _current_running_job_id
        running_proc = _current_process
        if not running_id or running_proc is None:
            return JSONResponse({"status": "idle", "detail": "Nenhum job em execucao."})
        _stop_requested_jobs.add(running_id)
    _terminate_process(running_proc)
    return JSONResponse({"status": "stopping", "job_id": running_id})


@app.post("/api/jobs/force/stop/{job_id}")
def api_force_stop_job(job_id: str):
    job = jobs.get(job_id)
    if job is None:
        return JSONResponse({"error": "Job nao encontrado."}, status_code=404)
    if job.status not in {"RUNNING", "QUEUED"}:
        return JSONResponse({"status": "ignored", "detail": "Job nao esta em execucao."})
    job.status = "STOPPED"
    job.error = "STOPPED: finalizado manualmente."
    job.finished_at = _now_ts()
    if job.email_status == "PENDING":
        job.email_status = "SKIPPED"
    if job.whatsapp_status == "PENDING":
        job.whatsapp_status = "SKIPPED"
    if job.drive_status == "PENDING":
        job.drive_status = "SKIPPED"
    return JSONResponse({"status": "stopped", "job_id": job_id})


@app.get("/api/status/{job_id}")
def api_status(job_id: str):
    job = jobs.get(job_id)
    if job is None:
        return JSONResponse({"error": "Job nao encontrado."}, status_code=404)
    return JSONResponse(_job_payload(job))


@app.get("/api/manual/latest")
def api_manual_latest():
    return JSONResponse({"latest_manual_job": _job_payload(_latest_manual_job())})


@app.get("/api/daily/latest")
def api_daily_latest():
    payload = _public_daily_schedule()
    payload["latest_daily_job"] = _job_payload(_latest_daily_job())
    return JSONResponse(payload)


@app.post("/api/daily/settings")
def api_daily_settings_post(payload: DailyScheduleSettingsPayload):
    if payload.extra_slots_enabled is not None:
        _set_all_extra_daily_slots(bool(payload.extra_slots_enabled))
    elif payload.enabled is not None and payload.run_time.strip():
        updated = _set_daily_slot_enabled(payload.run_time, bool(payload.enabled))
        if not updated and bool(payload.enabled):
            if _add_daily_custom_time(payload.run_time):
                updated = _set_daily_slot_enabled(payload.run_time, True)
        if not updated:
            return JSONResponse({"error": "Horario adicional invalido."}, status_code=400)
    else:
        return JSONResponse({"error": "Informe um horario adicional valido."}, status_code=400)
    _save_daily_schedule_settings()
    response = _public_daily_schedule()
    response["latest_daily_job"] = _job_payload(_latest_daily_job())
    return JSONResponse(response)


@app.get("/api/output/latest")
def api_output_latest():
    return JSONResponse({"latest_output_job": _job_payload(_latest_output_job())})


@app.post("/api/email/{job_id}/send")
def api_send_email(job_id: str, to: str | None = None):
    job = jobs.get(job_id)
    if not job:
        return JSONResponse({"error": "Job nao encontrado."}, status_code=404)
    if job.status != "DONE":
        return JSONResponse({"error": "Somente jobs concluidos podem ser enviados por e-mail."}, status_code=400)

    recipients_override = _parse_recipients(to) if (to or "").strip() else None
    if (to or "").strip() and not recipients_override:
        return JSONResponse({"error": "Nenhum e-mail valido informado no parametro 'to'."}, status_code=400)

    status, error = _send_output_email(job, recipients_override)
    job.email_status = status
    job.email_error = error
    if status == "SENT":
        job.email_sent_at = _now_ts()
        return JSONResponse({"status": "sent", "job_id": job.job_id})

    code = 500 if status == "FAILED" else 400
    return JSONResponse({"status": status.lower(), "job_id": job.job_id, "detail": error}, status_code=code)


@app.post("/api/email/send/latest")
def api_send_latest(to: str | None = None):
    latest = _latest_done_job()
    if not latest:
        return JSONResponse({"error": "Nenhum output concluido disponivel para envio."}, status_code=404)
    recipients_override = _parse_recipients(to) if (to or "").strip() else None
    status, error = _send_output_email(latest, recipients_override)
    latest.email_status = status
    latest.email_error = error
    if status == "SENT":
        latest.email_sent_at = _now_ts()
        return JSONResponse({"status": "sent", "job_id": latest.job_id, "mode": "latest"})
    code = 500 if status == "FAILED" else 400
    return JSONResponse({"status": status.lower(), "job_id": latest.job_id, "detail": error}, status_code=code)


@app.post("/api/whatsapp/{job_id}/send")
def api_send_whatsapp(job_id: str, to: str | None = None):
    job = jobs.get(job_id)
    if not job:
        return JSONResponse({"error": "Job nao encontrado."}, status_code=404)
    if job.status != "DONE":
        return JSONResponse({"error": "Somente jobs concluidos podem ser enviados por WhatsApp."}, status_code=400)

    recipients_override = _parse_whatsapp_recipients(to) if (to or "").strip() else None
    if (to or "").strip() and not recipients_override:
        return JSONResponse({"error": "Nenhum numero WhatsApp valido informado no parametro 'to'."}, status_code=400)

    status, error = _send_output_whatsapp(job, recipients_override)
    job.whatsapp_status = status
    job.whatsapp_error = error
    if status == "SENT":
        job.whatsapp_sent_at = _now_ts()
        return JSONResponse({"status": "sent", "job_id": job.job_id})

    code = 500 if status == "FAILED" else 400
    return JSONResponse({"status": status.lower(), "job_id": job.job_id, "detail": error}, status_code=code)


@app.post("/api/whatsapp/send/latest")
def api_send_whatsapp_latest(to: str | None = None):
    latest = _latest_done_job()
    if not latest:
        return JSONResponse({"error": "Nenhum output concluido disponivel para envio."}, status_code=404)
    recipients_override = _parse_whatsapp_recipients(to) if (to or "").strip() else None
    status, error = _send_output_whatsapp(latest, recipients_override)
    latest.whatsapp_status = status
    latest.whatsapp_error = error
    if status == "SENT":
        latest.whatsapp_sent_at = _now_ts()
        return JSONResponse({"status": "sent", "job_id": latest.job_id, "mode": "latest"})
    code = 500 if status == "FAILED" else 400
    return JSONResponse({"status": status.lower(), "job_id": latest.job_id, "detail": error}, status_code=code)


@app.post("/api/drive/{job_id}/upload")
def api_upload_drive(job_id: str, folder_id: str | None = None, credentials_file: str | None = None):
    job = jobs.get(job_id)
    if not job:
        return JSONResponse({"error": "Job nao encontrado."}, status_code=404)
    if job.status != "DONE":
        return JSONResponse({"error": "Somente jobs concluidos podem ser enviados para o Google Drive."}, status_code=400)

    folder_override = None
    if (folder_id or "").strip():
        folder_override = _normalize_drive_folder_id(folder_id)
        if not folder_override:
            return JSONResponse({"error": "Folder ID do Google Drive invalido."}, status_code=400)
    credentials_override = _normalize_drive_credentials_file(credentials_file) if (credentials_file or "").strip() else None

    status, error, file_id, file_url = _send_output_drive(job, folder_override, credentials_override)
    job.drive_status = status
    job.drive_error = error
    job.drive_file_id = file_id
    job.drive_file_url = file_url
    if folder_override:
        job.drive_folder_id = folder_override
    if credentials_override:
        job.drive_credentials_file = credentials_override
    if status == "SENT":
        job.drive_uploaded_at = _now_ts()
        return JSONResponse({"status": "sent", "job_id": job.job_id, "file_id": file_id, "file_url": file_url})

    code = 500 if status == "FAILED" else 400
    return JSONResponse({"status": status.lower(), "job_id": job.job_id, "detail": error}, status_code=code)


@app.post("/api/drive/upload/latest")
def api_upload_drive_latest(folder_id: str | None = None, credentials_file: str | None = None):
    latest = _latest_done_job()
    if not latest:
        return JSONResponse({"error": "Nenhum output concluido disponivel para upload."}, status_code=404)

    folder_override = None
    if (folder_id or "").strip():
        folder_override = _normalize_drive_folder_id(folder_id)
        if not folder_override:
            return JSONResponse({"error": "Folder ID do Google Drive invalido."}, status_code=400)
    credentials_override = _normalize_drive_credentials_file(credentials_file) if (credentials_file or "").strip() else None

    status, error, file_id, file_url = _send_output_drive(latest, folder_override, credentials_override)
    latest.drive_status = status
    latest.drive_error = error
    latest.drive_file_id = file_id
    latest.drive_file_url = file_url
    if folder_override:
        latest.drive_folder_id = folder_override
    if credentials_override:
        latest.drive_credentials_file = credentials_override
    if status == "SENT":
        latest.drive_uploaded_at = _now_ts()
        return JSONResponse(
            {"status": "sent", "job_id": latest.job_id, "mode": "latest", "file_id": file_id, "file_url": file_url}
        )

    code = 500 if status == "FAILED" else 400
    return JSONResponse({"status": status.lower(), "job_id": latest.job_id, "detail": error}, status_code=code)


@app.get("/api/events")
async def api_events(request: Request, job_id: str | None = None):
    async def stream():
        last = ""
        while True:
            if await request.is_disconnected():
                break
            current = jobs.get(job_id) if job_id else _latest_manual_job()
            payload = {
                "counts": _counts(),
                "latest_manual_job": _job_payload(_latest_manual_job()),
                "current_job": _job_payload(current),
            }
            encoded = json.dumps(payload, ensure_ascii=False)
            if encoded != last:
                yield f"event: snapshot\ndata: {encoded}\n\n"
                last = encoded
            await asyncio.sleep(2)

    return StreamingResponse(stream(), media_type="text/event-stream", headers={"Cache-Control": "no-cache"})


@app.get("/api/marketplace/credentials")
def api_marketplace_credentials(request: Request):
    return JSONResponse({"magalu": _magalu_payload(request)})


@app.post("/api/marketplace/credentials/magalu")
def api_marketplace_credentials_magalu(payload: MagaluCredentialsPayload, request: Request):
    updates = payload.model_dump(exclude_none=True)
    if not updates:
        return JSONResponse({"error": "Nenhum campo informado para atualizar."}, status_code=400)
    with _marketplace_lock:
        magalu = _marketplace_credentials.setdefault("magalu", _default_magalu_credentials())
        for key, value in updates.items():
            if key in magalu:
                magalu[key] = str(value).strip()
        _save_marketplace_credentials()
    return JSONResponse({"provider": "magalu", "updated_fields": sorted(updates.keys()), "credentials": _magalu_payload(request)})


@app.post("/api/admin/users/authenticate")
def api_admin_authenticate(payload: AdminAuthenticatePayload):
    username = str(payload.username or "").strip().lower()
    password = str(payload.password or "")
    with _admin_users_lock:
        user = _admin_users.get(username)
    if not user:
        return JSONResponse({"error": "Credenciais invalidas."}, status_code=401)
    if not user.get("is_active"):
        return JSONResponse({"error": "Usuario inativo."}, status_code=403)

    expected = str(user.get("password_hash") or "")
    salt = str(user.get("salt") or "")
    if not expected or not salt or not secrets.compare_digest(_hash_password(password, salt), expected):
        return JSONResponse({"error": "Credenciais invalidas."}, status_code=401)

    token = secrets.token_urlsafe(32)
    expires_in = USER_SESSION_TTL_SECONDS
    with _admin_sessions_lock:
        _admin_sessions[token] = {"username": username, "expires_at": _now_ts() + expires_in}

    return JSONResponse({"session_token": token, "expires_in_seconds": expires_in, "user": _public_user(user)})


@app.post("/api/admin/users/logout")
def api_admin_logout(request: Request):
    token = str(request.headers.get(ADMIN_SESSION_HEADER) or "").strip()
    if token:
        with _admin_sessions_lock:
            _admin_sessions.pop(token, None)
    return JSONResponse({"status": "ok"})


@app.get("/api/admin/users")
def api_admin_users(request: Request):
    actor, error = _admin_actor(request, require_admin=True)
    if error:
        return error
    with _admin_users_lock:
        users = [_public_user(user) for _, user in sorted(_admin_users.items(), key=lambda item: item[0])]
    return JSONResponse({"actor": _public_user(actor), "count": len(users), "users": users})


@app.post("/api/admin/users")
def api_admin_create_user(request: Request, payload: AdminCreateUserPayload):
    actor, error = _admin_actor(request, require_admin=True)
    if error:
        return error
    username = str(payload.username or "").strip().lower()
    password = str(payload.password or "")
    role = str(payload.role or "operator").strip().lower()

    if not re.fullmatch(r"[a-z0-9_.-]{3,64}", username):
        return JSONResponse({"error": "Username invalido. Use 3-64 chars [a-z0-9_.-]."}, status_code=400)
    if len(password) < 6:
        return JSONResponse({"error": "Senha deve ter ao menos 6 caracteres."}, status_code=400)
    if role not in {"admin", "operator"}:
        return JSONResponse({"error": "Perfil invalido. Use admin ou operator."}, status_code=400)

    with _admin_users_lock:
        if username in _admin_users:
            return JSONResponse({"error": "Usuario ja existe."}, status_code=409)
        hash_value, salt_hex = _new_password(password)
        now = _now_ts()
        _admin_users[username] = {
            "username": username,
            "full_name": str(payload.full_name or "").strip() or None,
            "role": role,
            "is_active": bool(payload.is_active),
            "password_hash": hash_value,
            "salt": salt_hex,
            "created_at": now,
            "updated_at": now,
        }
        _write_json(USERS_FILE, _admin_users)
        created = dict(_admin_users[username])

    return JSONResponse({"status": "created", "actor": _public_user(actor), "user": _public_user(created)})


@app.put("/api/admin/users/{username}")
def api_admin_update_user(request: Request, username: str, payload: AdminUpdateUserPayload):
    actor, error = _admin_actor(request, require_admin=True)
    if error:
        return error

    target = str(username or "").strip().lower()
    with _admin_users_lock:
        user = _admin_users.get(target)
        if not user:
            return JSONResponse({"error": "Usuario nao encontrado."}, status_code=404)

        if payload.role is not None:
            role = str(payload.role).strip().lower()
            if role not in {"admin", "operator"}:
                return JSONResponse({"error": "Perfil invalido. Use admin ou operator."}, status_code=400)
            user["role"] = role
        if payload.full_name is not None:
            user["full_name"] = str(payload.full_name).strip() or None
        if payload.is_active is not None:
            if target == "admin" and payload.is_active is False:
                return JSONResponse({"error": "Nao e permitido desativar o admin padrao."}, status_code=400)
            user["is_active"] = bool(payload.is_active)
        if payload.new_password is not None:
            if len(payload.new_password) < 6:
                return JSONResponse({"error": "Nova senha deve ter ao menos 6 caracteres."}, status_code=400)
            hash_value, salt_hex = _new_password(payload.new_password)
            user["password_hash"] = hash_value
            user["salt"] = salt_hex

        user["updated_at"] = _now_ts()
        _admin_users[target] = user
        _write_json(USERS_FILE, _admin_users)
        updated = dict(user)

    return JSONResponse({"status": "updated", "actor": _public_user(actor), "user": _public_user(updated)})


@app.delete("/api/admin/users/{username}")
def api_admin_delete_user(request: Request, username: str):
    actor, error = _admin_actor(request, require_admin=True)
    if error:
        return error

    target = str(username or "").strip().lower()
    if target == "admin":
        return JSONResponse({"error": "Nao e permitido excluir o admin padrao."}, status_code=400)
    if target == str(actor.get("username") or "").lower():
        return JSONResponse({"error": "Nao e permitido excluir o proprio usuario da sessao atual."}, status_code=400)

    with _admin_users_lock:
        if target not in _admin_users:
            return JSONResponse({"error": "Usuario nao encontrado."}, status_code=404)
        _admin_users.pop(target, None)
        _write_json(USERS_FILE, _admin_users)

    with _admin_sessions_lock:
        stale = [token for token, payload in _admin_sessions.items() if payload.get("username") == target]
        for token in stale:
            _admin_sessions.pop(token, None)

    return JSONResponse({"status": "deleted", "username": target})


@app.get("/api/functions/check")
def api_functions_check():
    checks = []
    checks.append({"name": "api_token", "status": "ok" if API_TOKEN else "warn", "detail": "API_TOKEN configurado." if API_TOKEN else "API_TOKEN vazio."})
    checks.append({"name": "default_input", "status": "ok" if DEFAULT_INPUT.exists() else "warn", "detail": "input.xlsx encontrado." if DEFAULT_INPUT.exists() else "input.xlsx nao encontrado."})
    checks.append({"name": "smtp", "status": "ok" if (SMTP_HOST and SMTP_SENDER) else "warn", "detail": "SMTP pronto." if (SMTP_HOST and SMTP_SENDER) else "SMTP nao configurado."})
    drive_settings = _public_drive_settings()
    checks.append(
        {
            "name": "google_drive",
            "status": "ok" if drive_settings["drive_ready"] else "warn",
            "detail": (
                "Google Drive pronto."
                if drive_settings["drive_ready"]
                else "Google Drive nao configurado (credenciais ou folder id ausentes)."
            ),
        }
    )

    summary = {"ok": 0, "warn": 0, "error": 0}
    for item in checks:
        summary[str(item["status"]).lower()] += 1

    status = "ok"
    if summary["error"] > 0:
        status = "error"
    elif summary["warn"] > 0:
        status = "warn"
    return JSONResponse({"status": status, "summary": summary, "checks": checks})


@app.get("/", response_class=HTMLResponse)
def index(request: Request):
    return templates.TemplateResponse(
        "index.html",
        {
            "request": request,
            "counts": _counts(),
            "has_default_input": DEFAULT_INPUT.exists(),
            "api_token": API_TOKEN,
            "assets_version": _assets_version(),
        },
    )


@app.get("/status/{job_id}", response_class=HTMLResponse)
def status(request: Request, job_id: str):
    job = jobs.get(job_id)
    if job is None:
        return HTMLResponse("Job nao encontrado.", status_code=404)
    return templates.TemplateResponse("status.html", {"request": request, "job": job, "assets_version": _assets_version()})


@app.get("/results/{job_id}", response_class=HTMLResponse)
def results(request: Request, job_id: str):
    job = jobs.get(job_id)
    if job is None:
        return HTMLResponse("Job nao encontrado.", status_code=404)
    if not job.output_path or not job.output_path.exists():
        return templates.TemplateResponse("status.html", {"request": request, "job": job, "assets_version": _assets_version()})
    headers, rows = read_output(job.output_path)
    return templates.TemplateResponse(
        "results.html",
        {"request": request, "headers": headers, "rows": rows, "job_id": job_id, "assets_version": _assets_version()},
    )


@app.get("/download/latest")
def download_latest(merged: int | None = None, mode: str | None = None):
    latest = _latest_output_job()
    if latest is None or not latest.output_path or not latest.output_path.exists():
        return HTMLResponse("Arquivo nao encontrado.", status_code=404)
    base_path = latest.output_path
    if merged and latest.email_extra_attachment:
        merged_path, _ = _build_merged_output(latest)
        if merged_path and merged_path.exists():
            base_path = merged_path
    if mode and mode.lower().strip() == "aprazo":
        simple_path, _ = _build_simple_output(base_path, "a_prazo")
        if simple_path and simple_path.exists():
            return FileResponse(
                path=simple_path,
                filename=simple_path.name,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
    return FileResponse(
        path=base_path,
        filename=base_path.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/download/daily/fixed")
def download_daily_fixed():
    fixed = _daily_result_path()
    if fixed.exists() and fixed.is_file():
        return FileResponse(
            path=fixed,
            filename=_daily_result_filename(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    latest = _latest_xlsx_in_dir(DAILY_RESULT_DIR)
    if latest is None or not latest.exists() or not latest.is_file():
        return HTMLResponse("Arquivo nao encontrado.", status_code=404)
    return FileResponse(
        path=latest,
        filename=_daily_result_filename(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/download/input-example")
def download_input_example():
    if not FIXED_MODEL_TEMPLATE.exists() or not FIXED_MODEL_TEMPLATE.is_file():
        return HTMLResponse("Arquivo modelo.xlsx nao encontrado.", status_code=404)
    return FileResponse(
        path=FIXED_MODEL_TEMPLATE,
        filename="Busca_Preco_modelo.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/download/{job_id}")
def download(job_id: str, merged: int | None = None, mode: str | None = None):
    job = jobs.get(job_id)
    if job is None or not job.output_path or not job.output_path.exists():
        return HTMLResponse("Arquivo nao encontrado.", status_code=404)
    base_path = job.output_path
    if merged and job.email_extra_attachment:
        merged_path, _ = _build_merged_output(job)
        if merged_path and merged_path.exists():
            base_path = merged_path
    if mode and mode.lower().strip() == "aprazo":
        simple_path, _ = _build_simple_output(base_path, "a_prazo")
        if simple_path and simple_path.exists():
            return FileResponse(
                path=simple_path,
                filename=simple_path.name,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
    return FileResponse(
        path=base_path,
        filename=base_path.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
