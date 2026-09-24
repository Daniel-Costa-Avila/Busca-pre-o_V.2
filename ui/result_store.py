"""Persistencia local dos resultados publicados pela API.

O banco e acessado somente pelo processo da API. Outros sistemas devem usar os
endpoints HTTP, nunca compartilhar o arquivo SQLite.
"""
from __future__ import annotations

import hashlib
import json
import sqlite3
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


RETENTION_DAYS = 7


def _utc_now() -> datetime:
    return datetime.now(timezone.utc)


def _connect(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(str(db_path), timeout=30)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    connection.execute("PRAGMA journal_mode = WAL")
    return connection


@contextmanager
def _database(db_path: Path):
    """Abre, confirma e fecha explicitamente a conexao (necessario no Windows)."""
    connection = _connect(db_path)
    try:
        yield connection
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def initialize(db_path: Path) -> None:
    with _database(db_path) as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS result_collections (
                id INTEGER PRIMARY KEY,
                job_id TEXT NOT NULL UNIQUE,
                source TEXT NOT NULL,
                finished_at TEXT NOT NULL,
                imported_at TEXT NOT NULL,
                total_items INTEGER NOT NULL,
                file_name TEXT NOT NULL,
                file_sha256 TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS result_items (
                id INTEGER PRIMARY KEY,
                collection_id INTEGER NOT NULL REFERENCES result_collections(id) ON DELETE CASCADE,
                row_number INTEGER NOT NULL,
                channel_id TEXT,
                internal_code TEXT,
                channel TEXT,
                title TEXT,
                price TEXT,
                link TEXT,
                payload_json TEXT NOT NULL
            );

            CREATE INDEX IF NOT EXISTS idx_result_collections_finished_at
                ON result_collections(finished_at DESC);
            CREATE INDEX IF NOT EXISTS idx_result_items_collection_id
                ON result_items(collection_id);
            CREATE INDEX IF NOT EXISTS idx_result_items_internal_code
                ON result_items(internal_code);
            CREATE INDEX IF NOT EXISTS idx_result_items_channel
                ON result_items(channel);
            CREATE INDEX IF NOT EXISTS idx_result_items_channel_id
                ON result_items(channel_id);
            """
        )


def _normalized_header(value: Any) -> str:
    return " ".join(str(value or "").casefold().replace("_", " ").split())


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _field(row: dict[str, Any], *names: str) -> str | None:
    normalized = {_normalized_header(key): value for key, value in row.items()}
    for name in names:
        value = normalized.get(_normalized_header(name))
        if value is not None:
            return _text(value)
    return None


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _read_rows(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["Output"] if "Output" in workbook.sheetnames else workbook.active
        headers = [str(cell.value or "").strip() for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
        rows: list[dict[str, Any]] = []
        for values in worksheet.iter_rows(min_row=2, values_only=True):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            rows.append({headers[index]: values[index] if index < len(values) else None for index in range(len(headers)) if headers[index]})
        return rows
    finally:
        workbook.close()


def purge_expired(db_path: Path, now: datetime | None = None) -> int:
    initialize(db_path)
    cutoff = ((now or _utc_now()) - timedelta(days=RETENTION_DAYS)).isoformat()
    with _database(db_path) as conn:
        cursor = conn.execute("DELETE FROM result_collections WHERE finished_at < ?", (cutoff,))
        return cursor.rowcount


def import_workbook(db_path: Path, job_id: str, source: str, output_path: Path, finished_at: datetime | None = None) -> int:
    """Importa um workbook uma unica vez por job e devolve sua collection id."""
    if not output_path.exists():
        raise FileNotFoundError(f"Arquivo de resultado nao encontrado: {output_path}")
    initialize(db_path)
    rows = _read_rows(output_path)
    completed_at = (finished_at or _utc_now()).astimezone(timezone.utc).isoformat()
    imported_at = _utc_now().isoformat()
    with _database(db_path) as conn:
        existing = conn.execute("SELECT id FROM result_collections WHERE job_id = ?", (job_id,)).fetchone()
        if existing:
            return int(existing["id"])
        cursor = conn.execute(
            """INSERT INTO result_collections
               (job_id, source, finished_at, imported_at, total_items, file_name, file_sha256)
               VALUES (?, ?, ?, ?, ?, ?, ?)""",
            (job_id, source, completed_at, imported_at, len(rows), output_path.name, _sha256(output_path)),
        )
        collection_id = int(cursor.lastrowid)
        conn.executemany(
            """INSERT INTO result_items
               (collection_id, row_number, channel_id, internal_code, channel, title, price, link, payload_json)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            [
                (
                    collection_id,
                    row_number,
                    _field(row, "id no canal", "id canal", "id"),
                    _field(row, "codigo interno", "código interno", "sku"),
                    _field(row, "canal"),
                    _field(row, "titulo", "título", "title"),
                    _field(row, "preco", "preço", "price"),
                    _field(row, "link", "url"),
                    json.dumps(row, ensure_ascii=False, default=str),
                )
                for row_number, row in enumerate(rows, start=2)
            ],
        )
    purge_expired(db_path)
    return collection_id


def _collection_payload(row: sqlite3.Row) -> dict[str, Any]:
    return dict(row)


def latest_collection(db_path: Path) -> dict[str, Any] | None:
    initialize(db_path)
    with _database(db_path) as conn:
        row = conn.execute("SELECT * FROM result_collections ORDER BY finished_at DESC LIMIT 1").fetchone()
        return _collection_payload(row) if row else None


def list_collections(db_path: Path, days: int) -> list[dict[str, Any]]:
    initialize(db_path)
    cutoff = (_utc_now() - timedelta(days=days)).isoformat()
    with _database(db_path) as conn:
        rows = conn.execute(
            "SELECT * FROM result_collections WHERE finished_at >= ? ORDER BY finished_at DESC", (cutoff,)
        ).fetchall()
        return [_collection_payload(row) for row in rows]


def query_items(
    db_path: Path, days: int, channel: str | None, internal_code: str | None,
    channel_id: str | None, query: str | None, page: int, page_size: int,
) -> tuple[int, list[dict[str, Any]]]:
    initialize(db_path)
    clauses = ["c.finished_at >= ?"]
    params: list[Any] = [(_utc_now() - timedelta(days=days)).isoformat()]
    for column, value in (("i.channel", channel), ("i.internal_code", internal_code), ("i.channel_id", channel_id)):
        if value:
            clauses.append(f"{column} = ?")
            params.append(value.strip())
    if query:
        clauses.append("(i.title LIKE ? OR i.internal_code LIKE ? OR i.channel_id LIKE ?)")
        params.extend([f"%{query.strip()}%"] * 3)
    where = " AND ".join(clauses)
    with _database(db_path) as conn:
        total = int(conn.execute(f"SELECT COUNT(*) FROM result_items i JOIN result_collections c ON c.id = i.collection_id WHERE {where}", params).fetchone()[0])
        rows = conn.execute(
            f"""SELECT i.row_number, i.channel_id, i.internal_code, i.channel, i.title, i.price, i.link,
                       i.payload_json, c.id AS collection_id, c.job_id, c.source, c.finished_at
                FROM result_items i JOIN result_collections c ON c.id = i.collection_id
                WHERE {where} ORDER BY c.finished_at DESC, i.row_number ASC LIMIT ? OFFSET ?""",
            [*params, page_size, (page - 1) * page_size],
        ).fetchall()
    items = []
    for row in rows:
        item = dict(row)
        item["payload"] = json.loads(item.pop("payload_json"))
        items.append(item)
    return total, items


def summary(db_path: Path, days: int) -> dict[str, Any]:
    collections = list_collections(db_path, days)
    with _database(db_path) as conn:
        cutoff = (_utc_now() - timedelta(days=days)).isoformat()
        total_items = int(conn.execute(
            "SELECT COUNT(*) FROM result_items i JOIN result_collections c ON c.id = i.collection_id WHERE c.finished_at >= ?",
            (cutoff,),
        ).fetchone()[0])
        channels = [dict(row) for row in conn.execute(
            """SELECT i.channel, COUNT(*) AS total FROM result_items i JOIN result_collections c ON c.id = i.collection_id
               WHERE c.finished_at >= ? GROUP BY i.channel ORDER BY total DESC""", (cutoff,)
        ).fetchall()]
    return {"days": days, "collections": len(collections), "items": total_items, "by_channel": channels}
