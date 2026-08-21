from __future__ import annotations

import re
import unicodedata
from collections.abc import Iterable
from copy import copy
from math import isfinite

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


HEADER_FILL = PatternFill("solid", fgColor="3B5A86")
ROW_FILL_LIGHT = PatternFill("solid", fgColor="E8EEF7")
ROW_FILL_BLUE = PatternFill("solid", fgColor="D6E1F0")
LOWEST_PRICE_FILL = PatternFill("solid", fgColor="C6EFCE")
WHITE_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=11, color="172033")
THIN_BLUE = Side(style="thin", color="AFC0D8")
CELL_BORDER = Border(left=THIN_BLUE, right=THIN_BLUE, top=THIN_BLUE, bottom=THIN_BLUE)
CURRENCY_FORMAT = '#,##0.00'
INTEGER_FORMAT = '#,##0'
PRICE_COMPARISON_HEADERS = {
    "probel (oficial)",
    "magazine luiza",
    "casas bahia",
    "web continental",
    "madeiramadeira",
    "zema",
    "mercado livre principal",
    "mercado livre cotia",
    "site probel",
}
REMOVED_RESULT_HEADERS = {
    "casa e video",
    "carrefour",
}
RENAMED_RESULT_HEADERS = {
    "mercado livre": "Mercado Livre Principal",
}


def _place_column_before(ws: Worksheet, source_header: str, target_header: str) -> None:
    headers = [ws.cell(row=1, column=column).value for column in range(1, ws.max_column + 1)]
    normalized_headers = [_normalize(header) for header in headers]
    normalized_source = _normalize(source_header)
    normalized_target = _normalize(target_header)

    if normalized_source not in normalized_headers or normalized_target not in normalized_headers:
        return

    source_index = normalized_headers.index(normalized_source)
    target_index = normalized_headers.index(normalized_target)
    if source_index + 1 == target_index:
        return

    column_order = list(range(ws.max_column))
    moved_column = column_order.pop(source_index)
    target_index = column_order.index(target_index)
    column_order.insert(target_index, moved_column)

    for row in range(1, ws.max_row + 1):
        snapshots = []
        for column in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=column)
            snapshots.append(
                (
                    cell.value,
                    copy(cell._style),
                    copy(cell.hyperlink),
                    copy(cell.comment),
                )
            )

        for destination, source in enumerate(column_order, start=1):
            value, cell_style, hyperlink, comment = snapshots[source]
            cell = ws.cell(row=row, column=destination)
            cell.value = value
            cell._style = cell_style
            cell._hyperlink = hyperlink
            cell.comment = comment


def ensure_result_column_schema(ws: Worksheet) -> bool:
    changed = False
    for column in range(ws.max_column, 0, -1):
        normalized_header = _normalize(ws.cell(row=1, column=column).value)
        if normalized_header in REMOVED_RESULT_HEADERS:
            ws.delete_cols(column, 1)
            changed = True

    for column in range(1, ws.max_column + 1):
        normalized_header = _normalize(ws.cell(row=1, column=column).value)
        renamed_header = RENAMED_RESULT_HEADERS.get(normalized_header)
        if renamed_header:
            ws.cell(row=1, column=column).value = renamed_header
            changed = True

    normalized_headers = [
        _normalize(ws.cell(row=1, column=column).value)
        for column in range(1, ws.max_column + 1)
    ]
    ml_principal = _normalize("Mercado Livre Principal")
    ml_cotia = _normalize("Mercado Livre Cotia")
    if ml_principal in normalized_headers and ml_cotia not in normalized_headers:
        principal_column = normalized_headers.index(ml_principal) + 1
        ws.insert_cols(principal_column + 1, 1)
        ws.cell(row=1, column=principal_column + 1).value = "Mercado Livre Cotia"
        changed = True

    return changed


def _recalculate_summary_metrics(ws: Worksheet) -> None:
    headers = [ws.cell(row=1, column=column).value for column in range(1, ws.max_column + 1)]
    normalized_headers = [_normalize(header) for header in headers]
    metric_names = {
        "store": "loja menor preco",
        "seller": "seller menor preco",
        "minimum": "menor preco",
        "average": "preco medio",
        "count": "quantidade de lojas",
    }
    if not all(name in normalized_headers for name in metric_names.values()):
        return

    metric_columns = {
        key: normalized_headers.index(name) + 1
        for key, name in metric_names.items()
    }
    comparison_columns = [
        column
        for column, header in enumerate(normalized_headers, start=1)
        if header in PRICE_COMPARISON_HEADERS
    ]

    for row in range(2, ws.max_row + 1):
        prices: list[tuple[int, float]] = []
        for column in comparison_columns:
            value = ws.cell(row=row, column=column).value
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                continue
            price = round(float(value), 2)
            if isfinite(price) and price > 0:
                prices.append((column, price))

        if not prices:
            ws.cell(row=row, column=metric_columns["store"]).value = None
            ws.cell(row=row, column=metric_columns["seller"]).value = None
            ws.cell(row=row, column=metric_columns["minimum"]).value = None
            ws.cell(row=row, column=metric_columns["average"]).value = None
            ws.cell(row=row, column=metric_columns["count"]).value = 0
            continue

        winner_column, minimum = min(prices, key=lambda item: item[1])
        winner_name = str(headers[winner_column - 1] or "").strip()
        store_cell = ws.cell(row=row, column=metric_columns["store"])
        seller_cell = ws.cell(row=row, column=metric_columns["seller"])
        previous_store = RENAMED_RESULT_HEADERS.get(
            _normalize(store_cell.value),
            str(store_cell.value or "").strip(),
        )
        previous_seller = str(seller_cell.value or "").strip()

        store_cell.value = winner_name
        seller_cell.value = (
            previous_seller
            if _normalize(previous_store) == _normalize(winner_name) and previous_seller
            else winner_name
        )
        ws.cell(row=row, column=metric_columns["minimum"]).value = minimum
        ws.cell(row=row, column=metric_columns["average"]).value = (
            sum(price for _, price in prices) / len(prices)
        )
        ws.cell(row=row, column=metric_columns["count"]).value = len(prices)


def ensure_result_column_order(ws: Worksheet) -> None:
    schema_changed = ensure_result_column_schema(ws)
    _place_column_before(ws, "Probel (oficial)", "Magazine Luiza")
    if schema_changed:
        _recalculate_summary_metrics(ws)


def _normalize(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return " ".join(text.strip().lower().split())


def _is_currency_header(header: object) -> bool:
    normalized = _normalize(header)
    if not normalized:
        return False
    exact_names = {
        "preco",
        "a prazo",
        "a vista",
        "menor preco",
        "preco medio",
        "probel (oficial)",
        "magazine luiza",
        "casas bahia",
        "web continental",
        "casa e video",
        "madeiramadeira",
        "zema",
        "mercado livre principal",
        "mercado livre cotia",
        "site probel",
    }
    if normalized in exact_names:
        return True
    return bool(re.search(r"(^|\s)preco($|\s)", normalized)) and "quantidade" not in normalized


def _is_integer_header(header: object) -> bool:
    normalized = _normalize(header)
    return normalized in {"quantidade de lojas", "quantidade", "qtd", "total"}


def _lowest_price_columns(
    ws: Worksheet,
    row: int,
    comparison_columns: set[int],
) -> set[int]:
    prices: list[tuple[int, float]] = []
    for column in comparison_columns:
        value = ws.cell(row=row, column=column).value
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            continue
        normalized_price = round(float(value), 2)
        if isfinite(normalized_price) and normalized_price > 0:
            prices.append((column, normalized_price))

    if not prices:
        return set()

    lowest_price = min(price for _, price in prices)
    return {column for column, price in prices if price == lowest_price}


def _column_width(header: object, values: Iterable[object]) -> float:
    normalized = _normalize(header)
    if normalized in {"titulo", "produto", "descricao"}:
        return 48.0
    if normalized in {"link", "url", "href"}:
        return 44.0
    if normalized in {"codigo interno", "codigo lojista", "id no canal", "sku"}:
        return 18.0
    if normalized in {"canal", "loja menor preco", "seller menor preco"}:
        return 23.0
    if _is_currency_header(header):
        return 18.0

    lengths = [len(str(header or ""))]
    lengths.extend(len(str(value or "")) for value in values)
    return float(min(32, max(12, max(lengths, default=12) + 2)))


def style_result_worksheet(ws: Worksheet) -> None:
    """Aplica o padrão visual oficial das planilhas de resultado."""
    if ws.max_row < 1 or ws.max_column < 1:
        return

    ensure_result_column_order(ws)

    max_row = ws.max_row
    max_col = ws.max_column
    header_values = [ws.cell(row=1, column=column).value for column in range(1, max_col + 1)]

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.row_dimensions[1].height = 36

    for column, header in enumerate(header_values, start=1):
        cell = ws.cell(row=1, column=column)
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.border = CELL_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        sample_values = (
            ws.cell(row=row, column=column).value
            for row in range(2, min(max_row, 80) + 1)
        )
        ws.column_dimensions[get_column_letter(column)].width = _column_width(header, sample_values)

    currency_columns = {
        column
        for column, header in enumerate(header_values, start=1)
        if _is_currency_header(header)
    }
    integer_columns = {
        column
        for column, header in enumerate(header_values, start=1)
        if _is_integer_header(header)
    }
    comparison_columns = {
        column
        for column, header in enumerate(header_values, start=1)
        if _normalize(header) in PRICE_COMPARISON_HEADERS
    }

    for row in range(2, max_row + 1):
        fill = ROW_FILL_BLUE if row % 2 == 0 else ROW_FILL_LIGHT
        lowest_price_columns = _lowest_price_columns(ws, row, comparison_columns)
        ws.row_dimensions[row].height = 20
        for column in range(1, max_col + 1):
            cell = ws.cell(row=row, column=column)
            cell.fill = LOWEST_PRICE_FILL if column in lowest_price_columns else fill
            cell.font = BODY_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(
                horizontal="right" if column in currency_columns or column in integer_columns else "left",
                vertical="center",
                wrap_text=False,
            )
            if column in currency_columns and isinstance(cell.value, (int, float)):
                cell.number_format = CURRENCY_FORMAT
            elif column in integer_columns and isinstance(cell.value, (int, float)):
                cell.number_format = INTEGER_FORMAT


def style_result_workbook(workbook) -> None:
    for worksheet in workbook.worksheets:
        style_result_worksheet(worksheet)
