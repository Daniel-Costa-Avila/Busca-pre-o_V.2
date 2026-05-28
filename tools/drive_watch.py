import json
import os
import re
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path
import unicodedata

import openpyxl  # type: ignore
from google.oauth2 import service_account  # type: ignore
from googleapiclient.discovery import build  # type: ignore
from googleapiclient.http import MediaIoBaseDownload, MediaFileUpload  # type: ignore

BASE_DIR = Path(__file__).resolve().parent.parent
RUNS_DIR = BASE_DIR / "runs"
INBOX_DIR = RUNS_DIR / "drive_watch_inbox"
STATE_PATH = RUNS_DIR / "drive_watch_state.json"
MANUAL_ACTIVATION_DIR = Path(r"C:\Users\daniel.avila\Desktop\AGENTE_DE_PRECOS\Ativação Manual")
PRIMARY_INPUT_PATH = BASE_DIR / "input.xlsx"
SECONDARY_INPUT_PATH = Path(r"C:\Users\daniel.avila\Desktop\AGENTE_DE_PRECOS\Planilha diaria\Relatorio_Financeiro.xlsx")
RESULT_DAILY_DIR = Path(r"C:\Users\daniel.avila\Desktop\AGENTE_DE_PRECOS\Resultado Diario")

RESULT_SYSTEM_NAME = (os.getenv("RESULT_SYSTEM_NAME") or "Busca-Preço").strip() or "Busca-Preço"


def _slugify_filename(text: str) -> str:
    raw = str(text or "").strip()
    if not raw:
        return "arquivo"
    normalized = unicodedata.normalize("NFKD", raw)
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    normalized = normalized.replace(" ", "-")
    normalized = re.sub(r"[^A-Za-z0-9_-]+", "-", normalized)
    normalized = re.sub(r"[-_]{2,}", "-", normalized)
    normalized = normalized.strip("-_")
    return normalized or "arquivo"


def _daily_result_filename(dt: datetime | None = None) -> str:
    stamp = (dt or datetime.now()).strftime("%d%m%Y")
    system_slug = _slugify_filename(RESULT_SYSTEM_NAME)
    return f"{system_slug}_{stamp}.xlsx"


def _daily_result_path(dt: datetime | None = None) -> Path:
    return RESULT_DAILY_DIR / _daily_result_filename(dt)

def _load_env_file(path: Path) -> None:
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        raw = line.strip()
        if not raw or raw.startswith("#") or "=" not in raw:
            continue
        key, value = raw.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


_load_env_file(BASE_DIR / ".env")

DRIVE_FOLDER_ID = (os.getenv("DRIVE_WATCH_FOLDER_ID") or "").strip()
CREDENTIALS_FILE = (os.getenv("GOOGLE_DRIVE_CREDENTIALS_FILE") or "").strip()
POLL_SECONDS = int((os.getenv("DRIVE_WATCH_POLL_SECONDS") or "30").strip() or 30)
OUTPUT_DRIVE_FILENAME = (os.getenv("DRIVE_WATCH_OUTPUT_NAME") or _daily_result_filename()).strip()


def _load_state() -> dict:
    if not STATE_PATH.exists():
        return {"processed_ids": []}
    try:
        return json.loads(STATE_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {"processed_ids": []}


def _save_state(state: dict) -> None:
    STATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    STATE_PATH.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")


def _build_service():
    if not CREDENTIALS_FILE:
        raise RuntimeError("GOOGLE_DRIVE_CREDENTIALS_FILE nao configurado.")
    creds = service_account.Credentials.from_service_account_file(
        CREDENTIALS_FILE,
        scopes=["https://www.googleapis.com/auth/drive.readonly"],
    )
    return build("drive", "v3", credentials=creds, cache_discovery=False)


def _list_files(service):
    query = (
        f"'{DRIVE_FOLDER_ID}' in parents and trashed = false and "
        "mimeType = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'"
    )
    result = service.files().list(
        q=query,
        fields="files(id, name, modifiedTime)",
        orderBy="modifiedTime desc",
    ).execute()
    return result.get("files", [])


def _download_file(service, file_id: str, filename: str) -> Path:
    INBOX_DIR.mkdir(parents=True, exist_ok=True)
    target = INBOX_DIR / filename
    request = service.files().get_media(fileId=file_id)
    with target.open("wb") as fh:
        downloader = MediaIoBaseDownload(fh, request)
        done = False
        while not done:
            _, done = downloader.next_chunk()
    return target


def _run_job(input_path: Path) -> Path:
    job_id = datetime.now().strftime("%Y%m%d_%H%M%S")
    job_dir = RUNS_DIR / job_id
    job_dir.mkdir(parents=True, exist_ok=True)
    output_path = job_dir / f"pesquisa_{job_id}.xlsx"
    if not PRIMARY_INPUT_PATH.exists():
        raise RuntimeError(f"Planilha principal nao encontrada: {PRIMARY_INPUT_PATH}")
    cmd = [sys.executable, "-m", "App.main", "--input", str(PRIMARY_INPUT_PATH), "--output", str(output_path)]
    if SECONDARY_INPUT_PATH.exists():
        cmd.extend(["--input-ml", str(SECONDARY_INPUT_PATH)])
    else:
        raise RuntimeError(f"Planilha secundaria nao encontrada: {SECONDARY_INPUT_PATH}")
    result = subprocess.run(cmd, cwd=str(BASE_DIR))
    if result.returncode != 0:
        raise RuntimeError(f"App.main falhou (exit={result.returncode}).")
    daily_target = _daily_result_path()
    daily_target.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(["powershell", "-NoProfile", "-Command", f"Copy-Item -Path '{output_path}' -Destination '{daily_target}' -Force"])
    return daily_target


def _append_output_to_dir(output_path: Path) -> Path:
    if not output_path.exists():
        raise RuntimeError("Arquivo de output nao encontrado para append.")
    MANUAL_ACTIVATION_DIR.mkdir(parents=True, exist_ok=True)

    target = MANUAL_ACTIVATION_DIR / _daily_result_filename()

    wb_src = openpyxl.load_workbook(output_path)
    ws_src = wb_src.active

    if target.exists():
        wb_tgt = openpyxl.load_workbook(target)
        ws_tgt = wb_tgt.active
        start_row = 2
    else:
        wb_tgt = openpyxl.Workbook()
        ws_tgt = wb_tgt.active
        start_row = 1

    for row in ws_src.iter_rows(min_row=start_row, values_only=True):
        ws_tgt.append(list(row))

    wb_tgt.save(target)
    return target


def _upload_result(service, folder_id: str, output_path: Path) -> None:
    if not output_path.exists():
        return
    query = (
        f"'{folder_id}' in parents and trashed = false and name = '{OUTPUT_DRIVE_FILENAME}'"
    )
    existing = service.files().list(q=query, fields="files(id, name)").execute().get("files", [])
    media = MediaFileUpload(
        str(output_path),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        resumable=False,
    )
    if existing:
        file_id = existing[0]["id"]
        service.files().update(fileId=file_id, media_body=media).execute()
    else:
        service.files().create(
            body={"name": OUTPUT_DRIVE_FILENAME, "parents": [folder_id]},
            media_body=media,
            fields="id",
        ).execute()


def main():
    if not DRIVE_FOLDER_ID:
        raise RuntimeError("DRIVE_WATCH_FOLDER_ID nao configurado.")

    service = _build_service()
    state = _load_state()
    processed = set(state.get("processed_ids", []))

    print("Drive watch ativo. Monitorando pasta:", DRIVE_FOLDER_ID)

    while True:
        files = _list_files(service)
        for item in files:
            file_id = item.get("id")
            name = item.get("name") or "arquivo.xlsx"
            if not file_id or file_id in processed:
                continue

            print("Novo arquivo:", name)
            local_path = _download_file(service, file_id, name)
            output_path = _run_job(local_path)
            _append_output_to_dir(output_path)
            _upload_result(service, DRIVE_FOLDER_ID, output_path)

            processed.add(file_id)
            state["processed_ids"] = sorted(processed)
            _save_state(state)

        time.sleep(POLL_SECONDS)


if __name__ == "__main__":
    main()
